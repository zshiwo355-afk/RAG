from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_product_mapping_review as review


class FakeTextQuery:
    def __init__(self, query_string, query_params=None):
        self.query_string = query_string
        self.query_params = query_params or {}


class FakeSearchRequest:
    def __init__(self, table_name, size, output_fields, text):
        self.table_name = table_name
        self.size = size
        self.output_fields = output_fields
        self.text = text


class FakeModels:
    TextQuery = FakeTextQuery
    SearchRequest = FakeSearchRequest


class FakeResponse:
    def __init__(self, rows):
        self.body = {"result": [{"fields": row} for row in rows]}


class FakeClient:
    def __init__(self, rows):
        self.rows = rows
        self.requests = []

    def search(self, request):
        self.requests.append(request)
        query = request.text.query_string
        matched = []
        for row in self.rows:
            text = " ".join(str(row.get(key, "")) for key in ["product_name", "source_text"])
            if "师藏年鉴" in query and "师藏年鉴" in text:
                matched.append(row)
            elif "陈任远" in query and "陈任远" in text:
                matched.append(row)
        return FakeResponse(matched[: request.size])


def write_mapping(path: Path, brand: str, count: int) -> list[str]:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "source_name",
        "source_doc_id",
        "source_sha1",
        "local_product_id",
        "local_product_name",
        "local_spec",
        "local_manufacturer",
        "local_barcode",
        "canonical_product_id",
        "canonical_product_name",
        "canonical_spec",
        "match_method",
        "confidence",
        "status",
        "notes",
    ]
    ids = []
    with path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fields)
        writer.writeheader()
        for index in range(count):
            local_id = f"local-{brand}-{index}"
            ids.append(local_id)
            writer.writerow(
                {
                    "source_name": f"{brand}.xlsx",
                    "source_doc_id": f"source-{brand}",
                    "source_sha1": "sha",
                    "local_product_id": local_id,
                    "local_product_name": f"{brand}酒·纪念{index}",
                    "local_spec": "500ml",
                    "local_manufacturer": f"{brand}酒业",
                    "local_barcode": "",
                    "canonical_product_id": "",
                    "canonical_product_name": "",
                    "canonical_spec": "",
                    "match_method": "",
                    "confidence": "medium",
                    "status": "needs_review",
                    "notes": "",
                }
            )
    return ids


def write_candidates(path: Path, local_ids: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    candidates = []
    for local_id in local_ids:
        candidates.append(
            {
                "local_product_id": local_id,
                "candidate_online_product_ids": [f"online-{local_id}-1", f"online-{local_id}-2", f"online-{local_id}-3"],
                "candidate_online_product_names": ["候选1", "候选2", "候选3"],
                "candidate_online_specs": ["500ml", "500ml", "500ml"],
                "match_evidence": ["name_high:1.00", "spec_match:1.00"],
                "confidence": "medium",
            }
        )
    path.write_text(json.dumps({"candidates": candidates}, ensure_ascii=False), encoding="utf-8")


@pytest.mark.offline
def test_product_mapping_review_filters_needs_review_and_writes_49_rows(tmp_path: Path) -> None:
    dmz_ids = write_mapping(tmp_path / "config/product_identity_map.csv", "大民族", 21)
    srx_ids = write_mapping(tmp_path / "config/product_identity_map_srx.csv", "石荣霄", 28)
    write_candidates(tmp_path / "output/acceptance/dmz_product_mapping_candidates_full.json", dmz_ids)
    write_candidates(tmp_path / "output/acceptance/srx_product_mapping_candidates_full.json", srx_ids)

    result = review.build_outputs(tmp_path, tmp_path / "output/review")
    payload = result["payload"]

    assert payload["total"] == 49
    assert payload["brand_counts"]["大民族"] == 21
    assert payload["brand_counts"]["石荣霄"] == 28
    assert payload["excel_dropdown_added"] is True
    assert (tmp_path / "output/review/product_mapping_review_49.xlsx").exists()
    rows = list(csv.DictReader((tmp_path / "output/review/product_mapping_review_49.csv").open(encoding="utf-8-sig")))
    assert len(rows) == 49
    assert rows[0]["candidate_1_product_id"]
    assert rows[0]["candidate_2_product_id"]
    assert rows[0]["candidate_3_product_id"]
    assert review.REVIEW_COLUMNS == payload["columns"]


@pytest.mark.offline
def test_chinese_product_mapping_review_keeps_raw_mapping_and_dropdowns(tmp_path: Path) -> None:
    dmz_ids = write_mapping(tmp_path / "config/product_identity_map.csv", "大民族", 21)
    srx_ids = write_mapping(tmp_path / "config/product_identity_map_srx.csv", "石荣霄", 28)
    write_candidates(tmp_path / "output/acceptance/dmz_product_mapping_candidates_full.json", dmz_ids)
    write_candidates(tmp_path / "output/acceptance/srx_product_mapping_candidates_full.json", srx_ids)

    result = review.build_chinese_outputs(tmp_path, tmp_path / "output/review")
    payload = result["payload"]

    assert payload["total"] == 49
    assert payload["brand_counts"]["大民族"] == 21
    assert payload["brand_counts"]["石荣霄"] == 28
    assert payload["selected_action_mapping_zh_to_raw"]["确认候选1"] == "confirm_candidate_1"
    assert "selected_action_raw" in payload["columns_zh"]
    workbook = load_workbook(tmp_path / "output/review/product_mapping_review_49_中文.xlsx")
    sheet = workbook.active
    assert sheet.max_row - 1 == 49
    validations = list(sheet.data_validations.dataValidation)
    assert any("确认候选1" in str(validation.formula1) for validation in validations)
    assert any("是,否" in str(validation.formula1) for validation in validations)


@pytest.mark.offline
def test_core_search_terms_keep_parenthesized_version_signal() -> None:
    terms = review.generate_core_search_terms("大民族酒·师藏年鉴（陈任远版）", "大民族")

    assert "师藏年鉴（陈任远版）" in terms
    assert "师藏年鉴陈任远版" in terms
    assert "师藏年鉴（陈仁远版）" in terms
    assert "师藏年鉴陈仁远版" in terms
    assert any("陈任远版" in term for term in terms)


@pytest.mark.offline
def test_enhanced_candidate_ranking_prefers_direct_core_name_hit() -> None:
    rows = [
        {
            "id": "wrong",
            "doc_id": "wrong",
            "product_id": "prod-wrong",
            "product_name": "大民族酒·端午纪念",
            "spec": "500ml*6整箱",
            "source_text": "大民族酒·端午纪念 500ml*6整箱",
        },
        {
            "id": "right",
            "doc_id": "right",
            "product_id": "prod-right",
            "product_name": "大民族酒·师藏年鉴陈任远版",
            "spec": "500ml*6整箱",
            "source_text": "产品名称：大民族酒·师藏年鉴（陈任远版） 规格：500ml*6整箱",
        },
    ]
    result = review.enhanced_candidates_for_row(
        FakeClient(rows),
        FakeModels,
        "text_docs",
        {
            "brand": "大民族",
            "local_product_name": "大民族酒·师藏年鉴（陈任远版）",
            "local_spec": "500ml*6整箱",
            "local_manufacturer": "贵州省仁怀市民族酒业有限公司",
        },
        top_k=10,
    )

    assert result["candidate_online_product_ids"][0] == "prod-right"
    assert any("核心名完整命中" in item for item in result["candidate_evidence_by_product"]["prod-right"])


@pytest.mark.offline
def test_enhanced_review_files_can_write_49_rows_with_ten_candidates(tmp_path: Path) -> None:
    rows = []
    for index in range(49):
        row = {
            "brand": "大民族" if index < 21 else "石荣霄",
            "source_name": "source.xlsx",
            "source_doc_id": "source-doc",
            "local_product_id": f"local-{index}",
            "local_product_name": f"产品{index}",
            "local_spec": "500ml",
            "local_manufacturer": "酒厂",
            "local_barcode": "",
            "source_file_path": "source.xlsx",
            "source_excel_row_number": str(index + 2),
            "confidence": "medium",
            "reason_for_review": "needs_review",
            "search_terms": [f"产品{index}"],
            "version_signals": [],
            "candidate_generation_note": "已按产品核心名称命中线上 text_docs 候选。",
            "candidate_scores": {},
        }
        for number in range(1, 11):
            row[f"candidate_{number}_product_id"] = f"online-{index}-{number}"
            row[f"candidate_{number}_product_name"] = f"候选{number}"
            row[f"candidate_{number}_spec"] = "500ml"
            row[f"candidate_{number}_evidence"] = "核心名完整命中"
        rows.append(row)

    csv_path = tmp_path / "review.csv"
    xlsx_path = tmp_path / "review.xlsx"
    json_path = tmp_path / "review.json"
    review.write_enhanced_csv(csv_path, rows)
    dropdown_added = review.write_enhanced_xlsx(xlsx_path, rows)
    payload = review.write_enhanced_json(json_path, rows, dropdown_added)

    assert payload["total"] == 49
    assert payload["brand_counts"]["大民族"] == 21
    assert payload["brand_counts"]["石荣霄"] == 28
    assert payload["selected_action_mapping_zh_to_raw"]["确认候选10"] == "confirm_candidate_10"
    assert "来源文件路径" in payload["columns_zh"]
    assert "源Excel行号" in payload["columns_zh"]
    assert "候选生成说明" in payload["columns_zh"]
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    assert sheet.max_row - 1 == 49
    assert any("确认候选10" in str(validation.formula1) for validation in sheet.data_validations.dataValidation)


@pytest.mark.offline
def test_enhanced_candidate_fallback_is_marked_as_reference_only() -> None:
    result = review.enhanced_candidates_for_row(
        FakeClient([]),
        FakeModels,
        "text_docs",
        {
            "brand": "大民族",
            "local_product_name": "大民族酒·师藏年鉴（张怀仁版）",
            "local_spec": "500ml*6整箱",
            "local_manufacturer": "贵州省仁怀市民族酒业有限公司",
        },
        fallback_candidate={
            "candidate_online_product_ids": ["prod-fallback"],
            "candidate_online_product_names": ["大民族酒·建厂42周年（臻藏版）"],
            "candidate_online_specs": ["500ml*6整箱"],
        },
        top_k=10,
    )

    assert result["fallback_used"] is True
    assert result["direct_core_hit_count"] == 0
    assert "未按核心名称命中" in result["candidate_generation_note"]
    assert "旧候选仅供参考" in result["candidate_evidence_by_product"]["prod-fallback"][-1]
