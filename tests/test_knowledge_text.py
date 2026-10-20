import io
import json
import zipfile
from xml.etree import ElementTree as ET
from html import escape
from html.parser import HTMLParser

import pytest

from rag_app import knowledge_text as kt


pytestmark = pytest.mark.offline


def office(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return output.getvalue()


def docx(body, **files):
    return office({"word/document.xml": '<w:document xmlns:w="%s"><w:body>%s</w:body></w:document>' % (kt.W[1:-1], body), **files})


def codes(result):
    return {item["code"] for item in result["issues"]}


def test_cleaning_is_disclosed_stable_and_preserves_meaning_and_urls():
    source = ('# 未验证\r\n    不能将 0/16 写成通过；12.5 mg，2026-09-27。\n'
              '号码 13900000000 / 13900000001 / +86 13900000000\n'
              '联系 demo@example.test；password="short value"；Bearer abcdefghijklmnop\n'
              'https://example.test/13900000000?id=13900000000&%74oken=demo-secret\n'
              'https://user:demo-password@example.test/path')
    cleaned, counts = kt.clean_text(source)
    assert '# 未验证\n    不能将 0/16 写成通过；12.5 mg，2026-09-27。' in cleaned
    phones = cleaned.splitlines()[2].split(' / ')
    assert phones[0].removeprefix('号码 ') == phones[2] and phones[1] != phones[2]
    assert 'https://example.test/13900000000?id=13900000000' in cleaned
    for private in ('demo@example.test', 'short value', 'abcdefghijklmnop', 'demo-secret', 'demo-password'):
        assert private not in cleaned and private not in json.dumps(counts)
    assert counts['phone_candidate'] == 3
    assert kt.clean_text(source) == (cleaned, counts)
    assert kt.clean_text(cleaned) == (cleaned, {})
    email_url, _ = kt.clean_text('https://example.test?email=demo@example.test')
    assert email_url.startswith('https://example.test?email=') and 'demo@example.test' not in email_url


@pytest.mark.parametrize('suffix', ['md', 'txt', 'csv', 'json', 'yaml', 'yml', 'xml', 'html'])
def test_utf8_formats_are_literal_text_not_code(suffix):
    content = '# 不执行\n    <script>throw Error("unchanged")</script>\n0/16，不采用'
    result = kt.parse_text_file('test.' + suffix, content.encode())
    assert result['text'] == content and result['locations'] == [{'kind': 'lines', 'start_line': 1}]
    assert result['issues'] == []


@pytest.mark.parametrize('bom,encoding', [(b'\xff\xfe', 'utf-16-le'), (b'\xfe\xff', 'utf-16-be')])
def test_utf16_requires_explicit_bom_and_reports_encoding_fallback(bom, encoding):
    content = '# 编码示例\r\n    不采用 0/16，12.5 mg'
    result = kt.parse_text_file('example.txt', bom + content.encode(encoding))
    assert result['text'] == content.replace('\r\n', '\n')
    assert result['issues'] == [{'code': 'encoding_fallback', 'encoding': encoding}]
    with pytest.raises(ValueError, match='^(non_utf8_text|non_text_content)$'):
        kt.parse_text_file('example.txt', content.encode(encoding))
    with pytest.raises(ValueError, match='^invalid_text_encoding$'):
        kt.parse_text_file('example.txt', bom + b'x')


def test_unknown_legacy_encoding_and_utf32_are_not_guessed():
    with pytest.raises(ValueError, match='^non_utf8_text$'):
        kt.parse_text_file('example.txt', '示例文本'.encode('gb18030'))
    with pytest.raises(ValueError, match='^unsupported_text_encoding$'):
        kt.parse_text_file('example.txt', '示例文本'.encode('utf-32'))


def test_docx_keeps_paragraph_table_order_and_discloses_omissions():
    raw = docx('<w:p><w:r><w:t>前文</w:t><w:tab/><w:t>否</w:t></w:r></w:p>'
               '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>数量</w:t></w:r></w:p></w:tc>'
               '<w:tc><w:p><w:r><w:t>0/16</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
               '<w:p><w:r><w:t>后文</w:t></w:r></w:p>', **{
                   'word/header1.xml': '<header/>', 'word/vbaProject.bin': b'never execute',
                   'word/_rels/document.xml.rels': '<Relationships><Relationship TargetMode="External" Target="https://example.test/private"/></Relationships>',
               })
    result = kt.parse_text_file('test.docx', raw)
    assert result['text'].startswith('前文\t否\n[表格 1 行 1] ["数量", "0/16"]\n后文')
    assert [x['kind'] for x in result['locations']] == ['paragraph', 'table_row', 'paragraph']
    assert codes(result) == {'docx_auxiliary_parts_not_parsed', 'macros_not_executed', 'external_links_not_fetched'}
    assert 'private' not in json.dumps(result['issues'])


def paragraph(text, style=None):
    return '<w:p>' + ('<w:pPr><w:pStyle w:val="%s"/></w:pPr>' % style if style else '') + '<w:r><w:t>' + escape(text) + '</w:t></w:r></w:p>'


def table(rows):
    return '<w:tbl>' + ''.join('<w:tr>' + ''.join('<w:tc>' + paragraph(cell) + '</w:tc>' for cell in row) + '</w:tr>' for row in rows) + '</w:tbl>'


@pytest.mark.parametrize('part', ['footnotes', 'endnotes', 'comments'])
def test_docx_empty_auxiliary_parts_do_not_report_missing_content(part):
    children = ''
    if part != 'comments':
        children = ''.join(
            '<w:%s w:type="%s" w:id="%s"><w:p><w:pPr><w:spacing w:line="240" w:lineRule="auto"/></w:pPr>'
            '<w:r><w:%s/></w:r></w:p></w:%s>' % (part[:-1], kind, index, kind, part[:-1])
            for index, kind in enumerate(('separator', 'continuationSeparator')))
    auxiliary = '<w:%s xmlns:w="%s">%s</w:%s>' % (part, kt.W[1:-1], children, part)
    raw = docx(paragraph('岗位SOP') + table([['所属部门', '电商部', '岗位名称', '电商负责人']]),
               **{'word/' + part + '.xml': auxiliary})
    result = kt.parse_text_file('岗位SOP.docx', raw)
    assert result['issues'] == []
    assert result['sop_metadata'] == {'department': '电商部', 'role_name': '电商负责人'}
    assert 'docx_auxiliary_parts_not_parsed' not in kt.docx_html_preview(raw).decode()


@pytest.mark.parametrize('part,children', [
    ('footnotes', '<w:footnote w:id="2">' + paragraph('审批必须由负责人确认') + '</w:footnote>'),
    ('endnotes', '<w:endnote w:id="2">' + paragraph('实际尾注') + '</w:endnote>'),
    ('comments', '<w:comment w:id="0">' + paragraph('实际批注') + '</w:comment>'),
    ('comments', '<w:comment w:id="0"/>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/><w:t>真实说明</w:t></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/><w:drawing/></w:r></w:p></w:footnote>'),
    ('endnotes', '<w:endnote w:type="separator" w:id="0"><w:p><w:r><w:separator/><w:object/></w:r></w:p></w:endnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/><w:unknown/></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r/></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator unknown="value"/></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator" w:id="0"><w:p><w:pPr><w:unknown/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:type="separator"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'),
    ('footnotes', '<w:footnote w:id="2"><w:p/></w:footnote>'),
    ('comments', '<w:unknown/>'),
])
def test_docx_real_or_unknown_auxiliary_content_still_requires_review(part, children):
    auxiliary = '<w:%s xmlns:w="%s">%s</w:%s>' % (part, kt.W[1:-1], children, part)
    raw = docx(paragraph('正文'), **{'word/' + part + '.xml': auxiliary})
    assert 'docx_auxiliary_parts_not_parsed' in codes(kt.parse_text_file('test.docx', raw))
    assert 'docx_auxiliary_parts_not_parsed' in kt.docx_html_preview(raw).decode()


def test_docx_auxiliary_root_is_validated_before_ignoring_empty_parts():
    for auxiliary in ('<unknown/>', '<w:comments xmlns:w="%s" unknown="value"/>' % kt.W[1:-1]):
        raw = docx(paragraph('正文'), **{'word/comments.xml': auxiliary})
        assert 'docx_auxiliary_parts_not_parsed' in codes(kt.parse_text_file('test.docx', raw))


@pytest.mark.parametrize('extra,expected', [
    ([], {'department': '总经办', 'role_name': 'AI架构师'}),
    ([['所属部门', '总经办']], {'department': '总经办', 'role_name': 'AI架构师'}),
    ([['所属部门', '另一部门']], None),
    ([['所属公司及部门', '另一部门']], None),
])
def test_sop_department_alias_requires_consistent_explicit_values(extra, expected):
    raw = docx(paragraph('岗位SOP') + table([['所属公司及部门', '总经办', '岗位名称', 'AI架构师'], *extra]))
    assert kt.parse_text_file('岗位SOP.docx', raw).get('sop_metadata') == expected


def test_sop_headings_step_acceptance_pairs_and_auxiliary_text_are_retained():
    from rag_app.knowledge_chunking import split_structure_v2
    raw = docx(paragraph('测试岗位标准作业程序') +
               table([['字段', '内容', '字段', '内容'], ['所属部门', '总经办', '岗位名称', 'AI架构师']]) +
               paragraph('开发流程', 'Heading1') + paragraph('标准操作步骤', 'Heading3') +
               table([['步骤', '具体操作', '输出'], ['1', '先核对输入', '输入清单'], ['2', '随后修复', '变更']]) + paragraph('') +
               table([['步骤', '验收标准', '异常处理'], ['2', '回归成功', '停止发布'], ['1', '输入齐全', '补齐输入']]), **{
                   'word/header1.xml': '<w:hdr xmlns:w="%s">%s</w:hdr>' % (kt.W[1:-1], paragraph('内部使用')),
                   'word/footer1.xml': '<w:ftr xmlns:w="%s">%s</w:ftr>' % (kt.W[1:-1], paragraph('第 1 页')),
               })
    result = kt.parse_text_file('岗位SOP.docx', raw)
    assert result['sop_metadata'] == {'department': '总经办', 'role_name': 'AI架构师'}
    assert result['issues'] == []
    assert '| 1 | 先核对输入 | 输入清单 | 输入齐全 | 补齐输入 |' in result['text']
    assert '| 2 | 随后修复 | 变更 | 回归成功 | 停止发布 |' in result['text']
    assert '# 页眉\n内部使用' in result['text'] and '# 页脚\n第 1 页' in result['text']
    assert any(location.get('source_tables') == [2, 3] for location in result['locations'])
    chunks = split_structure_v2(result['text'])
    step_chunk = next(chunk for chunk in chunks if '先核对输入' in chunk['source_text'])
    assert step_chunk['section_path'] == ['开发流程', '标准操作步骤']
    assert '验收标准' in step_chunk['source_text'] and '输入齐全' in step_chunk['source_text']


def test_step_pairs_do_not_cross_sections_or_guess_missing_steps():
    operations = table([['步骤', '具体操作'], ['1', '甲'], ['2', '乙']])
    acceptance = table([['步骤', '验收标准'], ['1', '成功']])
    result = kt.parse_text_file('test.docx', docx(operations + acceptance))
    assert all(len(location.get('source_tables', [])) <= 1 for location in result['locations'])
    acceptance = table([['步骤', '验收标准'], ['1', '成功'], ['2', '成功']])
    result = kt.parse_text_file('test.docx', docx(operations + paragraph('下一流程', 'Heading1') + acceptance))
    assert all(len(location.get('source_tables', [])) <= 1 for location in result['locations'])


def test_sop_metadata_requires_explicit_consistent_fields_and_sop_title():
    fields = table([['所属部门', '总经办', '岗位名称', 'AI架构师']])
    assert 'sop_metadata' not in kt.parse_text_file('岗位SOP.docx', docx(paragraph('案例') + fields))
    conflicting = fields + table([['岗位名称', '另一岗位']])
    assert 'sop_metadata' not in kt.parse_text_file('sop.docx', docx(paragraph('岗位SOP') + conflicting))


def test_sop_metadata_document_name_is_explicit_and_unambiguous():
    fields = table([['所属部门', '总经办', '岗位名称', 'AI架构师'], ['文档名称', 'AI系统交付标准作业程序']])
    body = paragraph('岗位SOP') + fields + paragraph('1 文档基本信息', 'Heading1')
    result = kt.parse_text_file('sop.docx', docx(body))
    assert result['sop_metadata'] == {
        'department': '总经办', 'role_name': 'AI架构师', 'document_name': 'AI系统交付标准作业程序',
    }
    conflicting = body + table([['文档名称', '另一个名称']])
    assert kt.parse_text_file('sop.docx', docx(conflicting))['sop_metadata'] == {
        'department': '总经办', 'role_name': 'AI架构师',
    }


def test_docx_html_preview_escapes_all_document_text_and_does_not_change_bytes():
    import hashlib
    attack = '<script>alert(1)</script><img src="https://private.test" onerror="alert(2)">'
    raw = docx(paragraph('中文预览', 'Heading1') + paragraph(attack) + table([['字段', '说明'], ['输入', attack]]))
    before = hashlib.sha256(raw).hexdigest()
    preview = kt.docx_html_preview(raw, '</title><script>evil()</script>').decode('utf-8')
    assert '<h1>中文预览</h1>' in preview and '<table>' in preview
    assert '&lt;script&gt;' in preview and 'Content-Security-Policy' in preview
    assert hashlib.sha256(raw).hexdigest() == before
    class Tags(HTMLParser):
        def __init__(self):
            super().__init__()
            self.tags = []
        def handle_starttag(self, tag, attrs):
            self.tags.append(tag)
            assert not any(key.startswith('on') or key in {'src', 'href'} for key, value in attrs)
    parser = Tags()
    parser.feed(preview)
    assert not {'script', 'img', 'iframe', 'object', 'a'} & set(parser.tags)


def test_docx_preview_discloses_omissions_and_rejects_unsafe_xml():
    preview = kt.docx_html_preview(docx(paragraph('正文') + '<w:p><w:r><w:drawing/></w:r></w:p>')).decode()
    assert 'docx_nontext_content_not_parsed' in preview
    with pytest.raises(ValueError, match='unsafe_xml_declaration'):
        kt.docx_html_preview(office({'word/document.xml': '<!DOCTYPE doc []><doc/>'}))


def test_xlsx_keeps_headers_rows_formula_and_cached_value_without_recalculation():
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = '实测'
    book.active.append(['数量', '公式'])
    book.active.append([3, '=A2*2'])
    book.active['A2'].number_format = '0.0" mg"'
    book.create_sheet('边界').append(['不适用'])
    output = io.BytesIO()
    book.save(output)
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        files = {item.filename: archive.read(item) for item in archive.infolist()}
    sheet = ET.fromstring(files['xl/worksheets/sheet1.xml'])
    cached = sheet.find('.//s:c[@r="B2"]/s:v', {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'})
    assert cached is not None
    cached.text = '6'
    files['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
    result = kt.parse_text_file('test.xlsx', office(files))
    assert '# 工作表 1：实测' in result['text'] and '# 工作表 2：边界' in result['text']
    assert 'A1="数量"' in result['text'] and 'B2=formula="=A2*2"; cached_value=6' in result['text']
    assert 'A2=3; number_format="0.0\\" mg\\""' in result['text']
    assert codes(result) == {'xlsx_formulas_not_recalculated'}
    assert {'kind': 'sheet_row', 'sheet_index': 1, 'row': 2} in result['locations']
    missing = kt.parse_text_file('test.xlsx', output.getvalue())
    assert 'xlsx_formula_cache_missing' in codes(missing)


def test_pdf_blank_page_needs_ocr_and_encryption_is_rejected():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    result = kt.parse_text_file('blank.pdf', output.getvalue())
    assert result['text'] == '' and result['locations'] == [{'kind': 'page', 'page': 1}]
    assert {'code': 'needs_ocr', 'page': 1} in result['issues']
    writer.encrypt('demo-pass')
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match='^encrypted_document$'):
        kt.parse_text_file('locked.pdf', output.getvalue())


def test_limits_are_explicit_instead_of_silent_partial_success(monkeypatch):
    monkeypatch.setattr(kt, 'MAX_TEXT_CHARS', 8)
    result = kt.parse_text_file('test.txt', b'123456789')
    assert result['text'] == '12345678' and 'text_character_limit' in codes(result)
    monkeypatch.setattr(kt, 'MAX_LINES', 2)
    result = kt.parse_text_file('test.txt', b'a\nb\nc')
    assert result['text'] == 'a\nb' and 'text_line_limit' in codes(result)
    monkeypatch.setattr(kt, 'MAX_EXPANDED_BYTES', 8)
    with pytest.raises(ValueError, match='^office_expanded_size_limit$'):
        kt.parse_text_file('test.docx', docx('<w:p/>'))
    monkeypatch.setattr(kt, 'MAX_FILE_BYTES', 8)
    with pytest.raises(ValueError, match='^file_size_limit$'):
        kt.parse_text_file('test.txt', b'123456789')


def test_spreadsheet_limits_and_false_dimensions_are_disclosed(monkeypatch):
    from openpyxl import Workbook

    book = Workbook()
    book.active.append(['header', 'second'])
    book.active.append([2, 3])
    output = io.BytesIO()
    book.save(output)
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        files = {item.filename: archive.read(item) for item in archive.infolist()}
    files['xl/worksheets/sheet1.xml'] = files['xl/worksheets/sheet1.xml'].replace(b'ref="A1:B2"', b'ref="A1:A1"')
    raw = office(files)
    assert 'B2=3' in kt.parse_text_file('test.xlsx', raw)['text']
    monkeypatch.setattr(kt, 'MAX_ROWS', 1)
    assert 'xlsx_rows_or_cells_limit' in codes(kt.parse_text_file('test.xlsx', raw))
    monkeypatch.setattr(kt, 'MAX_ROWS', 20)
    monkeypatch.setattr(kt, 'MAX_COLUMNS', 1)
    assert 'xlsx_columns_limit' in codes(kt.parse_text_file('test.xlsx', raw))


@pytest.mark.parametrize('name,data,code', [
    ('auth.json', b'demo-private-value', 'credential_file_skipped'),
    ('run.py', b'raise Exception("demo-private-value")', 'unsupported_file_type'),
    ('bad.txt', b'\xffdemo-private-value', 'non_utf8_text'),
    ('bad.pdf', b'demo-private-value', 'parse_failed'),
    ('bad.docx', b'demo-private-value', 'parse_failed'),
    ('bad.xlsx', office({'a.xml': '<!DOCTYPE root [<!ENTITY secret "demo-private-value">]><root/>'}), 'unsafe_xml_declaration'),
    ('bad.xlsx', office({'a.xml': '<!DOCTYPE root [<!ENTITY secret "demo-private-value">]><root/>'.encode('utf-16')}), 'unsafe_xml_declaration'),
    ('bad.docx', office({'../outside.xml': 'demo-private-value'}), 'unsafe_office_archive'),
])
def test_untrusted_failures_never_echo_document_content(name, data, code):
    with pytest.raises(ValueError) as error:
        kt.parse_text_file(name, data)
    assert str(error.value) == code


def test_credentials_recognized_across_path_styles():
    for name in ['.env', '.env.production', r'a\credentials.json', 'keys/id_ed25519', 'keys/CERT.P12']:
        assert kt.credential_name(name)
    assert not kt.credential_name('docs/key-concepts.md')
