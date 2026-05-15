#!/usr/bin/env python3
"""Build minimal local records for product Excel sources.

This is a conservative adapter. It does not replace the legacy Excel -> cleaned
products -> enriched products -> build_documents pipeline. For source Excel
files that are not already represented in output/products_enriched.json, it
creates needs_mapping records so operators can see exactly what still requires
field mapping before embedding.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PRODUCT_EXTENSIONS = {".xlsx", ".xlsm"}


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", " ", safe_text(value)).strip()


def sha1_text(text: str, length: int = 16) -> str:
    import hashlib

    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:length]


def sha1_file(path: Path) -> str:
    digest = hashlib.sha1()
    with path.open("rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relative_path(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def resolve_source(root: Path, source: str) -> Path:
    path = Path(source)
    if not path.is_absolute():
        path = root / path
    return path.resolve()


def discover_excel_files(source_path: Path) -> list[Path]:
    if source_path.is_file():
        return [source_path] if source_path.suffix.lower() in PRODUCT_EXTENSIONS and not source_path.name.startswith("~$") else []
    if not source_path.exists():
        raise FileNotFoundError(f"source 不存在：{source_path}")
    return sorted(
        path
        for path in source_path.rglob("*")
        if path.is_file()
        and path.suffix.lower() in PRODUCT_EXTENSIONS
        and not path.name.startswith("~$")
    )


def load_json_list(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        return []
    return [item for item in payload if isinstance(item, dict)]


def load_products_by_source(path: Path) -> dict[str, list[dict[str, Any]]]:
    payload = load_json_list(path)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in payload:
        source_file = safe_text(item.get("source_file")).strip()
        if source_file:
            grouped.setdefault(Path(source_file).as_posix(), []).append(item)
    return grouped


def load_documents_by_source(path: Path) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in load_json_list(path):
        metadata = item.get("metadata") if isinstance(item.get("metadata"), dict) else {}
        source_file = safe_text(metadata.get("source_file") or metadata.get("source_path")).strip()
        if source_file:
            grouped.setdefault(Path(source_file).as_posix(), []).append(item)
    return grouped


def product_excel_source_doc_id(source_sha1: str, source_file: str, fallback: str = "") -> str:
    if source_sha1:
        return f"product_excel__{source_sha1}"
    return "product_excel__" + sha1_text(source_file or fallback, 40)


def product_source_doc_id(product_id: str, suffix: str = "") -> str:
    if not product_id:
        return ""
    return f"product__{product_id}__{suffix}" if suffix else f"product__{product_id}"


def normalize_product_document_source(document: dict[str, Any], source_file_sha1: str = "", source_file_override: str = "") -> dict[str, Any]:
    metadata = document.get("metadata") if isinstance(document.get("metadata"), dict) else {}
    normalized = dict(document)
    normalized_metadata = dict(metadata)
    product_id = safe_text(normalized_metadata.get("product_id")).strip()
    source_file = source_file_override or safe_text(normalized_metadata.get("source_file") or normalized_metadata.get("source_path")).strip()
    source_sha1 = source_file_sha1 or safe_text(normalized_metadata.get("source_file_sha1") or normalized_metadata.get("source_sha1")).strip()
    previous_source_doc_id = safe_text(normalized_metadata.get("source_doc_id")).strip()
    field_name = safe_text(normalized_metadata.get("field_name")).strip()
    doc_type = safe_text(normalized_metadata.get("doc_type")).strip()

    if previous_source_doc_id and not previous_source_doc_id.startswith("product_excel__"):
        normalized_metadata.setdefault("legacy_product_source_doc_id", previous_source_doc_id)
        if previous_source_doc_id.startswith("product__"):
            normalized_metadata.setdefault("product_source_doc_id", previous_source_doc_id)

    if product_id:
        normalized_metadata.setdefault(
            "product_source_doc_id",
            product_source_doc_id(product_id, "" if doc_type in {"product_full", "product_excel_source_row"} else field_name or "field"),
        )
        normalized_metadata.setdefault("product_doc_group_id", product_source_doc_id(product_id))

    normalized_metadata["source_doc_id"] = product_excel_source_doc_id(source_sha1, source_file, product_id)
    normalized_metadata["source_sha1"] = source_sha1
    normalized_metadata["source_file_sha1"] = source_sha1
    normalized_metadata["source_type"] = "product_excel"
    normalized_metadata.setdefault("source_kind", "product_excel")
    if source_file:
        normalized_metadata["source_file"] = source_file
        normalized_metadata["source_path"] = source_file
    normalized["metadata"] = normalized_metadata
    return normalized


def worksheet_headers(ws, max_scan_rows: int = 12) -> tuple[int | None, list[str]]:
    best_row = None
    best_values: list[str] = []
    best_score = 0
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_scan_rows), values_only=True), start=1):
        values = [compact_text(value) for value in row if compact_text(value)]
        score = sum(1 for value in values if any(keyword in value for keyword in ["产品", "名称", "品牌", "规格", "包装", "酒质", "卖点"]))
        if score > best_score:
            best_score = score
            best_row = row_index
            best_values = values
    return best_row, best_values


def build_needs_mapping_rows(root: Path, excel_path: Path, max_preview_rows: int) -> list[dict[str, Any]]:
    rel_path = relative_path(excel_path, root)
    workbook = load_workbook(excel_path, data_only=False, read_only=True)
    rows: list[dict[str, Any]] = []
    for ws in workbook.worksheets:
        header_row, headers = worksheet_headers(ws)
        non_empty_rows = 0
        previews: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            values = [compact_text(value) for value in row]
            if any(values):
                non_empty_rows += 1
                if len(previews) < max_preview_rows:
                    previews.append([value for value in values if value][:12])
        source_doc_id = "product_excel__" + sha1_text(f"{rel_path}::{ws.title}", length=24)
        rows.append(
            {
                "source_path": rel_path,
                "source_type": "product_excel",
                "source_sheet": ws.title,
                "source_doc_id": source_doc_id,
                "product_id": "unknown",
                "status": "needs_mapping",
                "header_row": header_row,
                "headers": headers,
                "row_count": max(non_empty_rows - (1 if header_row else 0), 0),
                "preview_rows": previews,
                "created_at": datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds"),
                "message": "Excel source is not yet represented in output/products_enriched.json; confirm field mapping before embedding.",
            }
        )
    workbook.close()
    return rows


def product_to_document(product: dict[str, Any]) -> dict[str, Any]:
    product_id = safe_text(product.get("product_id")).strip()
    product_name = compact_text(product.get("product_name"))
    source_file = safe_text(product.get("source_file")).strip()
    source_sheet = safe_text(product.get("source_sheet")).strip()
    source_row = safe_text(product.get("source_row")).strip()
    source_sha1 = safe_text(product.get("source_file_sha1") or product.get("source_sha1")).strip()
    doc_id = f"{product_id}__ingest_build_source_row" if product_id else "product_excel__" + sha1_text(f"{source_file}::{source_sheet}::{source_row}", 24)
    lines = []
    for label, key in [
        ("产品名称", "product_name"),
        ("品牌", "brand"),
        ("系列", "series"),
        ("规格", "spec"),
        ("档次", "grade"),
        ("渠道", "channel"),
        ("价格", "price"),
        ("包装特点", "packaging_desc"),
        ("酒质介绍", "quality_desc"),
        ("卖点与优势", "selling_points"),
        ("产品故事", "product_story"),
        ("目标人群", "target_users"),
        ("礼赠属性与适用场景", "gift_attributes"),
    ]:
        value = product.get(key)
        if value:
            lines.append(f"{label}：{compact_text(value)}")
    page_content = "\n".join(lines).strip() or product_name or f"{source_file} {source_sheet} {source_row}"
    source_doc_id = product_excel_source_doc_id(source_sha1, source_file, product_id)
    return {
        "page_content": page_content,
        "metadata": {
            "doc_id": doc_id,
            "product_id": product_id or "unknown",
            "product_name": product_name,
            "doc_type": "product_excel_source_row",
            "field_name": "source_row",
            "source_kind": "product_excel",
            "source_type": "product_excel",
            "source_file": source_file,
            "source_path": source_file,
            "source_sheet": source_sheet,
            "source_row": source_row,
            "source_doc_id": source_doc_id,
            "source_sha1": source_sha1,
            "source_file_sha1": source_sha1,
            "product_source_doc_id": product_source_doc_id(product_id),
            "product_doc_group_id": product_source_doc_id(product_id),
            "build_status": "ready_from_products_enriched",
        },
    }


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file_obj:
        for row in rows:
            file_obj.write(json.dumps(row, ensure_ascii=False) + "\n")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build local product Excel ingest docs or needs_mapping records.")
    parser.add_argument("--root", default=str(ROOT), help="Project root.")
    parser.add_argument("--source", action="append", required=True, help="Excel file or directory. May be passed multiple times.")
    parser.add_argument("--products", default="output/products_enriched.json", help="Existing enriched products JSON.")
    parser.add_argument("--documents", default="output/documents_preview_v2.json", help="Existing product documents JSON.")
    parser.add_argument("--output-dir", default="output/ingest_build/product_excel", help="Output directory.")
    parser.add_argument("--max-preview-rows", type=int, default=5, help="Preview rows per sheet for needs_mapping records.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    output_dir = (root / args.output_dir).resolve()
    products_by_source = load_products_by_source(root / args.products)
    documents_by_source = load_documents_by_source(root / args.documents)

    standard_documents: list[dict[str, Any]] = []
    needs_mapping: list[dict[str, Any]] = []
    for source in args.source:
        source_path = resolve_source(root, source)
        for excel_path in discover_excel_files(source_path):
            rel_path = relative_path(excel_path, root)
            file_sha1 = sha1_file(excel_path)
            existing_documents = documents_by_source.get(rel_path, [])
            products = products_by_source.get(rel_path, [])
            if existing_documents:
                standard_documents.extend(normalize_product_document_source(document, file_sha1, rel_path) for document in existing_documents)
            elif products:
                standard_documents.extend(product_to_document({**product, "source_file_sha1": file_sha1}) for product in products)
            else:
                needs_mapping.extend(build_needs_mapping_rows(root, excel_path, args.max_preview_rows))

    output_dir.mkdir(parents=True, exist_ok=True)
    docs_path = output_dir / "documents.json"
    needs_mapping_path = output_dir / "needs_mapping.jsonl"
    report_path = output_dir / "build_report.md"
    docs_path.write_text(json.dumps(standard_documents, ensure_ascii=False, indent=2), encoding="utf-8")
    write_jsonl(needs_mapping_path, needs_mapping)
    report_path.write_text(
        "\n".join(
            [
                "# Product Excel Ingest Build Report",
                "",
                f"- Standard documents: {len(standard_documents)}",
                f"- needs_mapping records: {len(needs_mapping)}",
                f"- Documents: `{docs_path}`",
                f"- Needs mapping: `{needs_mapping_path}`",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"Standard documents: {len(standard_documents)}")
    print(f"needs_mapping records: {len(needs_mapping)}")
    print(f"Documents: {docs_path}")
    print(f"Needs mapping: {needs_mapping_path}")
    print(f"Report: {report_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("[ERROR] Interrupted by user.", file=sys.stderr)
        raise SystemExit(130)
    except Exception as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        raise SystemExit(1)
