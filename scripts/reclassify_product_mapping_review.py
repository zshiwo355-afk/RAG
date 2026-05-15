#!/usr/bin/env python3
"""Reclassify manual product mapping review rows into new-product candidates."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from build_product_mapping_review import ACTION_OPTIONS_ZH_ENHANCED  # type: ignore  # noqa: E402


DEFAULT_REVIEW_JSON = "output/review/product_mapping_review_49_增强版.json"
DEFAULT_REVIEW_CSV = "output/review/product_mapping_review_49_增强版.csv"
DEFAULT_MAP_DMz = "config/product_identity_map.csv"
DEFAULT_MAP_SRX = "config/product_identity_map_srx.csv"
DEFAULT_REGISTRY = "config/product_catalog_registry.csv"
DEFAULT_OUTPUT_PREFIX = "output/review/product_mapping_review_49_reclassified"
DEFAULT_NEW_PRODUCT_PREFIX = "output/review/new_product_candidates"

NEW_COLUMNS = [
    "线上是否核心名称命中",
    "是否只有宽泛候选",
    "建议动作",
    "建议原因",
    "是否建议作为新产品",
    "suggested_new_product_id",
]


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def truthy_text(value: Any) -> bool:
    return safe_text(value).lower() in {"true", "1", "yes", "是"}


def load_review_rows(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload.get("rows") if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        return []
    return [dict(row) for row in rows if isinstance(row, dict)]


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return [dict(row) for row in csv.DictReader(file_obj)]


def existing_product_ids(*paths: Path) -> set[str]:
    ids: set[str] = set()
    for path in paths:
        for row in read_csv_rows(path):
            for key in ("product_id", "canonical_product_id", "registry_product_id"):
                value = safe_text(row.get(key))
                if value:
                    ids.add(value)
    return ids


def next_registry_ids(existing_ids: set[str], count: int) -> list[str]:
    max_number = 0
    for value in existing_ids:
        match = re.fullmatch(r"prd_(\d+)", value)
        if match:
            max_number = max(max_number, int(match.group(1)))
    return [f"prd_{number:06d}" for number in range(max_number + 1, max_number + count + 1)]


def row_text(row: dict[str, Any]) -> str:
    parts = [
        safe_text(row.get("候选生成说明")),
        safe_text(row.get("需要人工确认原因")),
        *[safe_text(row.get(f"候选{index}匹配依据")) for index in range(1, 11)],
    ]
    return " ".join(part for part in parts if part)


def has_core_name_hit(row: dict[str, Any]) -> bool:
    text = row_text(row)
    if "无核心名称命中" in text or "未按核心名称命中" in text:
        return False
    return "核心名完整命中" in text or "核心名称命中" in text


def only_broad_fallback(row: dict[str, Any]) -> bool:
    text = row_text(row)
    if "旧候选仅供参考" in text or "宽泛" in text or "弱名称匹配" in text:
        return True
    if not has_core_name_hit(row):
        return True
    return False


def reclassify_row(row: dict[str, Any], suggested_id: str) -> dict[str, Any]:
    core_hit = has_core_name_hit(row)
    broad_only = only_broad_fallback(row)
    result = dict(row)
    if not core_hit and broad_only:
        action = "new_product_candidate"
        reason = "线上无可靠核心名称命中，现有候选仅作宽泛参考，不建议硬映射到已有产品。"
        recommend_new = True
        result["我的选择"] = "作为新产品"
        result["selected_action_raw"] = "create_new_product"
        result["selected_canonical_product_id"] = suggested_id
        result["selected_canonical_product_name"] = safe_text(row.get("本地产品名称"))
    elif core_hit:
        action = "map_to_existing_or_review"
        reason = "存在核心名称命中，应继续核对候选或保留人工确认。"
        recommend_new = False
    else:
        action = "needs_review"
        reason = "候选证据不足，仍需人工确认。"
        recommend_new = False
    result.update(
        {
            "线上是否核心名称命中": "是" if core_hit else "否",
            "是否只有宽泛候选": "是" if broad_only else "否",
            "建议动作": action,
            "建议原因": reason,
            "是否建议作为新产品": "是" if recommend_new else "否",
            "suggested_new_product_id": suggested_id if recommend_new else "",
        }
    )
    return result


def reclassify_rows(rows: list[dict[str, Any]], suggested_ids: list[str]) -> list[dict[str, Any]]:
    output = []
    id_index = 0
    for row in rows:
        suggested_id = suggested_ids[id_index] if id_index < len(suggested_ids) else ""
        classified = reclassify_row(row, suggested_id)
        if classified.get("是否建议作为新产品") == "是":
            id_index += 1
        output.append(classified)
    return output


def new_product_candidates(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates = []
    for row in rows:
        if row.get("是否建议作为新产品") != "是":
            continue
        candidates.append(
            {
                "brand": safe_text(row.get("品牌")),
                "source_name": safe_text(row.get("来源名称")),
                "source_doc_id": safe_text(row.get("来源文件ID")),
                "source_sha1": safe_text(row.get("来源文件ID")).rsplit("__", 1)[-1],
                "source_path": safe_text(row.get("来源文件路径")),
                "excel_row": safe_text(row.get("源Excel行号")),
                "local_product_id": safe_text(row.get("本地产品ID")),
                "local_product_name": safe_text(row.get("本地产品名称")),
                "local_spec": safe_text(row.get("本地规格")),
                "local_manufacturer": safe_text(row.get("本地生产酒厂")),
                "suggested_new_product_id": safe_text(row.get("suggested_new_product_id")),
                "suggested_doc_types": ["product_full", "product_field"],
                "reason": safe_text(row.get("建议原因")),
            }
        )
    return candidates


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([row.get(key, "") for key in fieldnames])
    if "我的选择" in fieldnames:
        col = fieldnames.index("我的选择") + 1
        formula = '"' + ",".join(ACTION_OPTIONS_ZH_ENHANCED) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(2, col).coordinate}:{ws.cell(max(len(rows) + 1, 2), col).coordinate}")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 12), 40)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_md(path: Path, rows: list[dict[str, Any]], candidates: list[dict[str, Any]]) -> None:
    counts = {
        "new_product_candidate": sum(1 for row in rows if row.get("建议动作") == "new_product_candidate"),
        "needs_review": sum(1 for row in rows if row.get("建议动作") == "needs_review"),
        "map_to_existing_or_review": sum(1 for row in rows if row.get("建议动作") == "map_to_existing_or_review"),
    }
    lines = [
        "# Product Mapping Review Reclassified",
        "",
        f"- total: {len(rows)}",
        f"- new_product_candidate: {counts['new_product_candidate']}",
        f"- needs_review: {counts['needs_review']}",
        f"- map_to_existing_or_review: {counts['map_to_existing_or_review']}",
        "",
        "## New Product Candidates",
        "",
    ]
    for item in candidates:
        lines.append(f"- `{item['suggested_new_product_id']}` {item['brand']} / {item['local_product_name']} / row {item['excel_row']}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_new_product_plan(prefix: Path, candidates: list[dict[str, Any]]) -> None:
    prefix.parent.mkdir(parents=True, exist_ok=True)
    (prefix.with_suffix(".json")).write_text(json.dumps({"total": len(candidates), "candidates": candidates}, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(prefix.with_suffix(".csv"), candidates)
    lines = ["# New Product Candidates", "", f"- total: {len(candidates)}", ""]
    for item in candidates:
        lines.append(f"- `{item['suggested_new_product_id']}` {item['brand']} / {item['local_product_name']} / row {item['excel_row']}")
    prefix.with_suffix(".md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Reclassify product mapping review rows.")
    parser.add_argument("--review-json", default=DEFAULT_REVIEW_JSON)
    parser.add_argument("--review-csv", default=DEFAULT_REVIEW_CSV)
    parser.add_argument("--dmz-map", default=DEFAULT_MAP_DMz)
    parser.add_argument("--srx-map", default=DEFAULT_MAP_SRX)
    parser.add_argument("--registry", default=DEFAULT_REGISTRY)
    parser.add_argument("--output-prefix", default=DEFAULT_OUTPUT_PREFIX)
    parser.add_argument("--new-product-prefix", default=DEFAULT_NEW_PRODUCT_PREFIX)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    rows = load_review_rows(ROOT / args.review_json)
    existing_ids = existing_product_ids(ROOT / args.registry, ROOT / args.dmz_map, ROOT / args.srx_map)
    suggested_ids = next_registry_ids(existing_ids, len(rows))
    reclassified = reclassify_rows(rows, suggested_ids)
    candidates = new_product_candidates(reclassified)
    payload = {
        "created_at": now_iso(),
        "total": len(reclassified),
        "new_product_candidate_count": len(candidates),
        "needs_review_count": sum(1 for row in reclassified if row.get("建议动作") == "needs_review"),
        "map_to_existing_or_review_count": sum(1 for row in reclassified if row.get("建议动作") == "map_to_existing_or_review"),
        "rows": reclassified,
    }
    prefix = ROOT / args.output_prefix
    prefix.parent.mkdir(parents=True, exist_ok=True)
    prefix.with_suffix(".json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(prefix.with_suffix(".csv"), reclassified)
    write_xlsx(prefix.with_suffix(".xlsx"), reclassified)
    write_md(prefix.with_suffix(".md"), reclassified, candidates)
    write_new_product_plan(ROOT / args.new_product_prefix, candidates)
    print(json.dumps({key: payload[key] for key in ("total", "new_product_candidate_count", "needs_review_count", "map_to_existing_or_review_count")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
