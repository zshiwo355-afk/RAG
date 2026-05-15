#!/usr/bin/env python3
"""Build local product documents for review rows selected as new products."""

from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
for path in [ROOT / "scripts", ROOT / "src"]:
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from check_existing_product_overlap import compact, document_identity, load_documents, safe_text  # type: ignore  # noqa: E402


DEFAULT_REVIEW = "output/review/product_mapping_review_49_reclassified.xlsx"
DEFAULT_DMZ_DOCUMENTS = "output/ingest_build/product_excel/documents.json"
DEFAULT_SRX_DOCUMENTS = "output/ingest_build/product_excel_srx/documents.json"
DEFAULT_OUTPUT_DIR = "output/ingest_build/product_excel_new_products"
DEFAULT_ACCEPTANCE_MD = "output/acceptance/new_product_documents_acceptance.md"
DEFAULT_ACCEPTANCE_JSON = "output/acceptance/new_product_documents_acceptance.json"


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def resolve_path(root: Path, value: str) -> Path:
    path = Path(value)
    if path.is_absolute():
        return path.resolve()
    return (root / path).resolve()


def safe_slug(value: str) -> str:
    text = re.sub(r"[^\w.-]+", "_", safe_text(value)).strip("_")
    return text or "item"


def load_review_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [safe_text(value) for value in rows[0]]
        return [
            {headers[index]: safe_text(value) for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(safe_text(value) for value in row)
        ]
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("rows") if isinstance(payload, dict) else payload
        return [dict(row) for row in rows if isinstance(row, dict)] if isinstance(rows, list) else []
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return [dict(row) for row in csv.DictReader(file_obj)]


def is_new_product_row(row: dict[str, Any]) -> bool:
    return (
        safe_text(row.get("我的选择")) == "作为新产品"
        or safe_text(row.get("selected_action_raw")) == "create_new_product"
        or safe_text(row.get("建议动作")) == "new_product_candidate"
        or safe_text(row.get("是否建议作为新产品")) == "是"
    )


def review_new_products(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result = {}
    for row in rows:
        if not is_new_product_row(row):
            continue
        local_product_id = safe_text(row.get("本地产品ID"))
        canonical_id = safe_text(row.get("suggested_new_product_id") or row.get("selected_canonical_product_id"))
        if local_product_id and canonical_id:
            result[local_product_id] = dict(row)
    return result


def metadata_of(document: dict[str, Any]) -> dict[str, Any]:
    metadata = document.get("metadata")
    return metadata if isinstance(metadata, dict) else {}


def text_of(document: dict[str, Any]) -> str:
    return compact(document.get("page_content") or document.get("text") or document.get("source_text"))


def section_for(document: dict[str, Any]) -> str:
    metadata = metadata_of(document)
    return safe_slug(safe_text(metadata.get("field_name") or metadata.get("doc_type") or document.get("doc_type") or "section"))


def build_new_product_documents(
    documents: list[dict[str, Any]],
    review_rows_by_local: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    docs_by_local: dict[str, list[dict[str, Any]]] = {}
    for document in documents:
        identity = document_identity(document)
        local_product_id = identity.get("product_id")
        if local_product_id in review_rows_by_local:
            docs_by_local.setdefault(local_product_id, []).append(document)

    output_docs: list[dict[str, Any]] = []
    seen_doc_ids: set[str] = set()
    conflicts: list[str] = []
    for local_product_id, local_docs in docs_by_local.items():
        row = review_rows_by_local[local_product_id]
        canonical_product_id = safe_text(row.get("suggested_new_product_id") or row.get("selected_canonical_product_id"))
        for index, document in enumerate(local_docs, start=1):
            metadata = metadata_of(document)
            identity = document_identity(document)
            original_doc_type = safe_text(metadata.get("doc_type") or document.get("doc_type"))
            doc_type = original_doc_type if original_doc_type in {"product_full", "product_field"} else "product_field"
            source_sha1 = identity.get("source_sha1")
            doc_id = f"newprod__{canonical_product_id}__{source_sha1[:12]}__{safe_slug(local_product_id)}__{section_for(document)}__{index:03d}"
            if doc_id in seen_doc_ids:
                conflicts.append(doc_id)
                continue
            seen_doc_ids.add(doc_id)
            text = text_of(document)
            new_metadata = dict(metadata)
            new_metadata.update(
                {
                    "doc_id": doc_id,
                    "product_id": canonical_product_id,
                    "canonical_product_id": canonical_product_id,
                    "registry_product_id": canonical_product_id,
                    "local_product_id": local_product_id,
                    "legacy_local_product_id": local_product_id,
                    "doc_type": doc_type,
                    "source_type": "product_excel_new_product",
                    "source_doc_id": identity.get("source_doc_id"),
                    "source_sha1": source_sha1,
                    "source_path": identity.get("source_path"),
                    "source_file": identity.get("source_path"),
                    "product_name": safe_text(row.get("本地产品名称")) or identity.get("product_name"),
                    "spec": safe_text(row.get("本地规格")) or identity.get("spec"),
                    "manufacturer": safe_text(row.get("本地生产酒厂")) or identity.get("manufacturer"),
                    "brand": safe_text(row.get("品牌")) or identity.get("brand"),
                    "original_doc_id": safe_text(metadata.get("doc_id") or document.get("doc_id")),
                    "original_product_id": local_product_id,
                }
            )
            output_docs.append(
                {
                    "doc_id": doc_id,
                    "product_id": canonical_product_id,
                    "canonical_product_id": canonical_product_id,
                    "registry_product_id": canonical_product_id,
                    "local_product_id": local_product_id,
                    "doc_type": doc_type,
                    "source_type": "product_excel_new_product",
                    "source_doc_id": identity.get("source_doc_id"),
                    "source_sha1": source_sha1,
                    "source_path": identity.get("source_path"),
                    "product_name": new_metadata["product_name"],
                    "spec": new_metadata["spec"],
                    "manufacturer": new_metadata["manufacturer"],
                    "brand": new_metadata["brand"],
                    "text": text,
                    "page_content": text,
                    "metadata": new_metadata,
                }
            )
    stats = {
        "selected_new_product_count": len(review_rows_by_local),
        "local_products_with_documents": len(docs_by_local),
        "new_product_document_count": len(output_docs),
        "doc_id_conflicts": conflicts,
        "doc_type_counts": dict(Counter(doc["doc_type"] for doc in output_docs)),
        "missing_local_product_ids": sorted(set(review_rows_by_local) - set(docs_by_local)),
    }
    return output_docs, stats


def write_reports(output_dir: Path, acceptance_md: Path, acceptance_json: Path, docs: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    docs_path = output_dir / "documents.json"
    build_report = output_dir / "build_report.md"
    docs_path.write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding="utf-8")
    payload = {
        **stats,
        "documents_path": str(docs_path),
        "created_at": now_iso(),
        "opensearch_written": False,
        "embedding_called": False,
        "push_executed": False,
    }
    acceptance_json.parent.mkdir(parents=True, exist_ok=True)
    acceptance_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [
        "# New Product Documents Build Report",
        "",
        f"- selected new products: {stats['selected_new_product_count']}",
        f"- local products with documents: {stats['local_products_with_documents']}",
        f"- documents: {stats['new_product_document_count']}",
        f"- doc_id conflicts: {len(stats['doc_id_conflicts'])}",
        f"- missing local products: {len(stats['missing_local_product_ids'])}",
        f"- output: `{docs_path}`",
    ]
    text = "\n".join(lines) + "\n"
    build_report.write_text(text, encoding="utf-8")
    acceptance_md.write_text(text, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build local new product documents from review workbook.")
    parser.add_argument("--root", default=str(ROOT))
    parser.add_argument("--review", default=DEFAULT_REVIEW)
    parser.add_argument("--dmz-documents", default=DEFAULT_DMZ_DOCUMENTS)
    parser.add_argument("--srx-documents", default=DEFAULT_SRX_DOCUMENTS)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--report-output", default=DEFAULT_ACCEPTANCE_MD)
    parser.add_argument("--json-output", default=DEFAULT_ACCEPTANCE_JSON)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    review_rows = load_review_rows(resolve_path(root, args.review))
    selected = review_new_products(review_rows)
    documents = load_documents(resolve_path(root, args.dmz_documents)) + load_documents(resolve_path(root, args.srx_documents))
    new_docs, stats = build_new_product_documents(documents, selected)
    write_reports(resolve_path(root, args.output_dir), resolve_path(root, args.report_output), resolve_path(root, args.json_output), new_docs, stats)
    print(json.dumps(stats, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
