#!/usr/bin/env python3
"""Build a manual review workbook for needs_review product mappings."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from difflib import SequenceMatcher
import json
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


for path in [ROOT := Path(__file__).resolve().parents[1], ROOT / "src", ROOT / "scripts"]:
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from check_existing_product_overlap import extract_result_items, flatten_item  # type: ignore  # noqa: E402
from push_text_docs_to_opensearch import create_client, ensure_runtime_config, load_env  # type: ignore  # noqa: E402
from rag_app.retrieval_service import build_text_query_string  # noqa: E402


DEFAULT_OUTPUT_DIR = "output/review"
ACTION_OPTIONS = [
    "confirm_candidate_1",
    "confirm_candidate_2",
    "confirm_candidate_3",
    "create_new_product",
    "ignore",
    "need_more_info",
]
CONFIRMED_OPTIONS = ["是", "否"]
ACTION_OPTIONS_ZH = [
    "确认候选1",
    "确认候选2",
    "确认候选3",
    "作为新产品",
    "忽略",
    "信息不足，稍后再看",
]
ACTION_OPTIONS_ZH_ENHANCED = [
    *(f"确认候选{index}" for index in range(1, 11)),
    "手动填写线上产品ID",
    "作为新产品",
    "忽略",
    "信息不足，稍后再看",
]
ACTION_ZH_TO_RAW = {
    "确认候选1": "confirm_candidate_1",
    "确认候选2": "confirm_candidate_2",
    "确认候选3": "confirm_candidate_3",
    "作为新产品": "create_new_product",
    "忽略": "ignore",
    "信息不足，稍后再看": "need_more_info",
}
ACTION_ZH_TO_RAW_ENHANCED = {
    **{f"确认候选{index}": f"confirm_candidate_{index}" for index in range(1, 11)},
    "手动填写线上产品ID": "manual_product_id",
    "作为新产品": "create_new_product",
    "忽略": "ignore",
    "信息不足，稍后再看": "need_more_info",
}
REVIEW_COLUMNS = [
    "brand",
    "source_name",
    "source_doc_id",
    "local_product_id",
    "local_product_name",
    "local_spec",
    "local_manufacturer",
    "local_barcode",
    "candidate_1_product_id",
    "candidate_1_product_name",
    "candidate_1_spec",
    "candidate_1_evidence",
    "candidate_2_product_id",
    "candidate_2_product_name",
    "candidate_2_spec",
    "candidate_2_evidence",
    "candidate_3_product_id",
    "candidate_3_product_name",
    "candidate_3_spec",
    "candidate_3_evidence",
    "confidence",
    "reason_for_review",
    "selected_action",
    "selected_canonical_product_id",
    "selected_canonical_product_name",
    "reviewer_note",
    "是否确认",
]
REVIEW_COLUMNS_ZH = [
    "品牌",
    "来源名称",
    "来源文件ID",
    "本地产品ID",
    "本地产品名称",
    "本地规格",
    "本地生产酒厂",
    "本地条码",
    "候选1产品ID",
    "候选1产品名称",
    "候选1规格",
    "候选1匹配依据",
    "候选2产品ID",
    "候选2产品名称",
    "候选2规格",
    "候选2匹配依据",
    "候选3产品ID",
    "候选3产品名称",
    "候选3规格",
    "候选3匹配依据",
    "置信度",
    "需要人工确认原因",
    "我的选择",
    "选择对应的线上产品ID",
    "选择对应的线上产品名称",
    "是否确认",
    "备注",
    "selected_action_raw",
    "selected_canonical_product_id",
    "selected_canonical_product_name",
]
ENHANCED_BASE_COLUMNS_ZH = [
    "品牌",
    "来源名称",
    "来源文件路径",
    "源Excel行号",
    "来源文件ID",
    "本地产品ID",
    "本地产品名称",
    "本地规格",
    "本地生产酒厂",
    "本地条码",
]
ENHANCED_TAIL_COLUMNS_ZH = [
    "置信度",
    "需要人工确认原因",
    "核心搜索词",
    "版本强信号",
    "候选生成说明",
    "我的选择",
    "手动填写线上产品ID",
    "选择对应的线上产品ID",
    "选择对应的线上产品名称",
    "是否确认",
    "备注",
    "selected_action_raw",
    "selected_canonical_product_id",
    "selected_canonical_product_name",
]
REVIEW_COLUMNS_ZH_ENHANCED = [
    *ENHANCED_BASE_COLUMNS_ZH,
    *[
        value
        for index in range(1, 11)
        for value in (f"候选{index}产品ID", f"候选{index}产品名称", f"候选{index}规格", f"候选{index}匹配依据")
    ],
    *ENHANCED_TAIL_COLUMNS_ZH,
]
ZH_COLUMN_TO_RAW = {
    "品牌": "brand",
    "来源名称": "source_name",
    "来源文件ID": "source_doc_id",
    "本地产品ID": "local_product_id",
    "本地产品名称": "local_product_name",
    "本地规格": "local_spec",
    "本地生产酒厂": "local_manufacturer",
    "本地条码": "local_barcode",
    "候选1产品ID": "candidate_1_product_id",
    "候选1产品名称": "candidate_1_product_name",
    "候选1规格": "candidate_1_spec",
    "候选1匹配依据": "candidate_1_evidence",
    "候选2产品ID": "candidate_2_product_id",
    "候选2产品名称": "candidate_2_product_name",
    "候选2规格": "candidate_2_spec",
    "候选2匹配依据": "candidate_2_evidence",
    "候选3产品ID": "candidate_3_product_id",
    "候选3产品名称": "candidate_3_product_name",
    "候选3规格": "candidate_3_spec",
    "候选3匹配依据": "candidate_3_evidence",
    "置信度": "confidence",
    "需要人工确认原因": "reason_for_review",
    "选择对应的线上产品ID": "selected_canonical_product_id",
    "选择对应的线上产品名称": "selected_canonical_product_name",
    "是否确认": "是否确认",
    "备注": "reviewer_note",
    "selected_action_raw": "selected_action",
    "selected_canonical_product_id": "selected_canonical_product_id",
    "selected_canonical_product_name": "selected_canonical_product_name",
}


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        return [dict(row) for row in csv.DictReader(file_obj)]


def load_candidates(path: Path) -> dict[str, dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    candidates = payload.get("candidates") if isinstance(payload, dict) else []
    if not isinstance(candidates, list):
        candidates = []
    return {
        safe_text(item.get("local_product_id")): item
        for item in candidates
        if isinstance(item, dict) and safe_text(item.get("local_product_id"))
    }


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", " ", safe_text(value)).strip()


def normalize_for_match(value: Any) -> str:
    text = compact_text(value).lower()
    return re.sub(r"[\s·•\-_（）()\[\]【】《》<>:：,，。;；/\\]+", "", text)


def brand_prefixes(brand: str) -> list[str]:
    if brand == "大民族":
        return ["大民族酒", "大民族"]
    if brand == "石荣霄":
        return ["石荣霄酒", "百年石荣霄酒", "石荣霄"]
    return [brand]


def strip_brand_prefix(product_name: str, brand: str) -> str:
    text = compact_text(product_name)
    for prefix in sorted(brand_prefixes(brand), key=len, reverse=True):
        for sep in ["·", "•", " ", ""]:
            candidate = f"{prefix}{sep}"
            if text.startswith(candidate):
                return text[len(candidate) :].strip()
    return text


def remove_symbols_keep_text(value: str) -> str:
    return re.sub(r"[\s·•\-_（）()\[\]【】《》<>:：,，。;；/\\]+", "", compact_text(value))


def spaced_core(value: str) -> str:
    text = compact_text(value)
    text = re.sub(r"[·•（）()\[\]【】《》<>:：,，。;；/\\]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_version_signals(product_name: str) -> list[str]:
    text = compact_text(product_name)
    signals: list[str] = []
    for match in re.findall(r"[（(]([^）)]+)[）)]", text):
        value = compact_text(match)
        if value:
            signals.append(value)
    patterns = [
        r"\b[VA]\d{1,2}\b",
        r"精品\s*\d+\s*[号號]",
        r"藏\s*\d+",
        r"天字号\s*[VA]\d{1,2}",
        r"师藏年鉴[^\s，,。；;]*",
        r"[甲乙丙丁戊己庚辛壬癸]?[子丑寅卯辰巳午未申酉戌亥]?[鼠牛虎兔龙蛇马羊猴鸡狗猪]年",
        r"端午纪念",
    ]
    for pattern in patterns:
        for match in re.findall(pattern, text, flags=re.I):
            value = compact_text(match)
            if value:
                signals.append(value)
    unique: list[str] = []
    seen: set[str] = set()
    for signal in signals:
        key = normalize_for_match(signal)
        if key and key not in seen:
            seen.add(key)
            unique.append(signal)
    return unique


def generate_core_search_terms(product_name: str, brand: str) -> list[str]:
    original = compact_text(product_name)
    core = strip_brand_prefix(original, brand)
    core_no_symbols = remove_symbols_keep_text(core)
    brand_core = spaced_core(f"{brand} {core}")
    terms = [original, core, core_no_symbols, brand_core]
    for signal in extract_version_signals(original):
        terms.append(signal)
        if core_no_symbols and normalize_for_match(signal) not in normalize_for_match(core_no_symbols):
            terms.append(f"{core_no_symbols}{remove_symbols_keep_text(signal)}")
    # The source Excel uses both 陈任远 and 陈仁远 in different cells for the
    # same product. Keep both spellings as strong exact-name search variants.
    for term in list(terms):
        if "陈任远" in term:
            terms.append(term.replace("陈任远", "陈仁远"))
        if "陈仁远" in term:
            terms.append(term.replace("陈仁远", "陈任远"))
    unique: list[str] = []
    seen: set[str] = set()
    for term in terms:
        term = compact_text(term)
        key = term.lower()
        if key and key not in seen:
            seen.add(key)
            unique.append(term)
    return unique


def parse_metadata(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value.strip():
        try:
            payload = json.loads(value)
        except Exception:
            return {}
        return payload if isinstance(payload, dict) else {}
    return {}


def extract_label(text: str, labels: list[str]) -> str:
    for label in labels:
        match = re.search(rf"{re.escape(label)}[:：]\s*([^\n；;，,]{{1,90}})", text)
        if match:
            return compact_text(match.group(1))
    return ""


def online_identity(row: dict[str, Any]) -> dict[str, Any]:
    metadata = parse_metadata(row.get("metadata"))
    merged = dict(metadata)
    merged.update({key: value for key, value in row.items() if value not in (None, "")})
    text = safe_text(merged.get("source_text") or merged.get("page_content") or merged.get("text")).strip()
    return {
        "doc_id": compact_text(merged.get("doc_id") or merged.get("id")),
        "product_id": compact_text(merged.get("product_id")),
        "product_name": compact_text(merged.get("product_name")) or extract_label(text, ["产品名称", "商品名称", "样品名称", "名称"]),
        "spec": compact_text(merged.get("spec")) or extract_label(text, ["规格型号", "规格", "净含量"]),
        "manufacturer": compact_text(merged.get("manufacturer")) or extract_label(text, ["生产厂家", "生产酒厂", "酒厂", "委托单位"]),
        "brand": compact_text(merged.get("brand")) or extract_label(text, ["品牌"]),
        "source_text": text,
        "text_preview": re.sub(r"\s+", " ", text)[:360],
        "score": row.get("score"),
    }


def output_fields_for_search() -> list[str]:
    return [
        "id",
        "doc_id",
        "product_id",
        "product_name",
        "brand",
        "spec",
        "manufacturer",
        "doc_type",
        "source_text",
        "source_doc_id",
        "metadata",
    ]


def search_text_docs(client: Any, models_module: Any, table_names: list[str], field_name: str, term: str, top_k: int) -> list[dict[str, Any]]:
    text = models_module.TextQuery(
        query_string=build_text_query_string(field_name, term),
        query_params={"default_op": "AND"},
    )
    last_error: Exception | None = None
    for table_name in table_names:
        try:
            request = models_module.SearchRequest(
                table_name=table_name,
                size=top_k,
                output_fields=output_fields_for_search(),
                text=text,
            )
            response = client.search(request)
            rows = [online_identity(flatten_item(row)) for row in extract_result_items(response)]
            if rows:
                return rows
        except Exception as exc:
            last_error = exc
    if last_error:
        raise last_error
    return []


def resolve_text_table_names(config: dict[str, str]) -> list[str]:
    return [f"{config['instance_id']}_text_docs", "text_docs"]


def name_similarity(left: str, right: str) -> float:
    left_n = normalize_for_match(left)
    right_n = normalize_for_match(right)
    if not left_n or not right_n:
        return 0.0
    if left_n == right_n:
        return 1.0
    if left_n in right_n or right_n in left_n:
        return 0.92
    return SequenceMatcher(None, left_n, right_n).ratio()


def score_candidate(local_row: dict[str, str], candidate: dict[str, Any], search_terms: list[str], version_signals: list[str]) -> tuple[float, list[str]]:
    product_name = candidate.get("product_name", "")
    source_text = candidate.get("source_text", "")
    haystack = f"{product_name} {source_text}"
    haystack_n = normalize_for_match(haystack)
    core = search_terms[1] if len(search_terms) > 1 else local_row.get("local_product_name", "")
    core_no_symbols = remove_symbols_keep_text(core)
    score = 0.0
    evidence: list[str] = []
    if core_no_symbols and normalize_for_match(core_no_symbols) in haystack_n:
        score += 100
        evidence.append(f"核心名完整命中:{core_no_symbols}")
    if product_name:
        sim = name_similarity(core, product_name)
        score += sim * 45
        if sim >= 0.88:
            evidence.append(f"产品名高相似:{sim:.2f}")
    for signal in version_signals:
        if normalize_for_match(signal) and normalize_for_match(signal) in haystack_n:
            score += 25
            evidence.append(f"版本强信号命中:{signal}")
    brand = local_row.get("brand", "")
    if brand and brand in haystack:
        score += 12
        evidence.append(f"同品牌:{brand}")
    local_spec = compact_text(local_row.get("local_spec"))
    if local_spec and normalize_for_match(local_spec) and normalize_for_match(local_spec) in haystack_n:
        score += 8
        evidence.append("规格一致")
    manufacturer = compact_text(local_row.get("local_manufacturer"))
    manufacturer_head = manufacturer.split()[0] if manufacturer else ""
    if manufacturer_head and len(manufacturer_head) >= 3 and manufacturer_head in haystack:
        score += 5
        evidence.append("酒厂线索一致")
    if not evidence:
        score += max(name_similarity(local_row.get("local_product_name", ""), product_name), 0) * 15
        evidence.append("弱名称匹配")
    return score, evidence


def enhanced_candidates_for_row(
    client: Any,
    models_module: Any,
    table_names: list[str],
    row: dict[str, str],
    fallback_candidate: dict[str, Any] | None = None,
    top_k: int = 10,
) -> dict[str, Any]:
    search_terms = generate_core_search_terms(row.get("local_product_name", ""), row.get("brand", ""))
    version_signals = extract_version_signals(row.get("local_product_name", ""))
    found: dict[str, dict[str, Any]] = {}
    evidence_by_product: dict[str, list[str]] = {}
    fields = ["product_name", "source_text"]
    for term in search_terms[:8]:
        for field_name in fields:
            try:
                rows = search_text_docs(client, models_module, table_names, field_name, term, top_k=20)
            except Exception:
                if field_name == "product_name":
                    continue
                rows = []
            for candidate in rows:
                product_id = compact_text(candidate.get("product_id"))
                if not product_id:
                    continue
                current = found.get(product_id)
                if current is None:
                    found[product_id] = candidate
                elif not current.get("product_name") and candidate.get("product_name"):
                    found[product_id] = candidate
                evidence_by_product.setdefault(product_id, []).append(f"{field_name}命中:{term}")

    direct_core_hit_count = len(found)
    fallback_used = False
    if not found and fallback_candidate:
        fallback_used = True
        ids = fallback_candidate.get("candidate_online_product_ids") or []
        names = fallback_candidate.get("candidate_online_product_names") or []
        specs = fallback_candidate.get("candidate_online_specs") or []
        for index, product_id in enumerate(ids[:top_k]):
            if not product_id:
                continue
            found[product_id] = {
                "product_id": product_id,
                "product_name": names[index] if index < len(names) else "",
                "spec": specs[index] if index < len(specs) else "",
                "source_text": "",
                "text_preview": "",
            }
            evidence_by_product[product_id] = ["无核心名称命中；旧候选仅供参考，不建议直接确认"]

    scored = []
    for product_id, candidate in found.items():
        score, evidence = score_candidate(row, candidate, search_terms, version_signals)
        scored.append((score, product_id, candidate, [*evidence, *evidence_by_product.get(product_id, [])]))
    scored.sort(key=lambda item: item[0], reverse=True)
    top = scored[:top_k]
    return {
        "search_terms": search_terms,
        "version_signals": version_signals,
        "candidate_online_product_ids": [item[1] for item in top],
        "candidate_online_product_names": [compact_text(item[2].get("product_name")) for item in top],
        "candidate_online_specs": [compact_text(item[2].get("spec")) for item in top],
        "candidate_online_doc_ids": [compact_text(item[2].get("doc_id")) for item in top],
        "candidate_evidence_by_product": {item[1]: item[3][:8] for item in top},
        "candidate_scores": {item[1]: round(float(item[0]), 4) for item in top},
        "direct_core_hit_count": direct_core_hit_count,
        "fallback_used": fallback_used,
        "candidate_generation_note": (
            "已按产品核心名称命中线上 text_docs 候选。"
            if direct_core_hit_count
            else "线上 text_docs 未按核心名称命中；本地源 Excel 记录已保留，请手动核对线上 product_id，或判断为新产品。旧候选仅供参考。"
        ),
    }


def candidate_evidence(candidate: dict[str, Any], index: int) -> str:
    evidence = candidate.get("match_evidence")
    if isinstance(evidence, list) and evidence:
        # Current candidate reports store aggregate evidence. Keep it compact
        # rather than trying to imply per-candidate precision that is not there.
        return "; ".join(safe_text(item) for item in evidence[:6] if safe_text(item))
    return ""


def risk_types(row: dict[str, str], candidate: dict[str, Any] | None) -> list[str]:
    risks: list[str] = []
    name = safe_text(row.get("local_product_name"))
    spec = safe_text(row.get("local_spec"))
    if not spec:
        risks.append("规格缺失")
    if re.search(r"(纪念|年份|生肖|端午|中秋|周年|限定|典藏|珍藏|藏\d+|V\d+|A\d+)", name):
        risks.append("纪念版/年份版/生肖版")
    candidate_names = [safe_text(item) for item in ((candidate or {}).get("candidate_online_product_names") or [])[:3]]
    candidate_specs = [safe_text(item) for item in ((candidate or {}).get("candidate_online_specs") or [])[:3]]
    if len({item for item in candidate_names if item}) > 1:
        risks.append("多个候选相近")
    if spec and sum(1 for item in candidate_specs if item and item != spec) >= 2:
        risks.append("同名多规格")
    brand = "大民族" if "大民族" in safe_text(row.get("source_name")) else "石荣霄"
    if any(item and brand not in item for item in candidate_names):
        risks.append("候选跨品牌")
    return risks or ["人工确认"]


def build_review_rows(
    mapping_specs: list[tuple[str, Path, Path]],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for brand, mapping_path, candidates_path in mapping_specs:
        mappings = [row for row in read_csv(mapping_path) if safe_text(row.get("status")) == "needs_review"]
        candidates_by_local_id = load_candidates(candidates_path)
        for row in mappings:
            candidate = candidates_by_local_id.get(safe_text(row.get("local_product_id"))) or {}
            product_ids = candidate.get("candidate_online_product_ids") if isinstance(candidate, dict) else []
            product_names = candidate.get("candidate_online_product_names") if isinstance(candidate, dict) else []
            specs = candidate.get("candidate_online_specs") if isinstance(candidate, dict) else []
            product_ids = product_ids if isinstance(product_ids, list) else []
            product_names = product_names if isinstance(product_names, list) else []
            specs = specs if isinstance(specs, list) else []
            item: dict[str, str] = {
                "brand": brand,
                "source_name": safe_text(row.get("source_name")),
                "source_doc_id": safe_text(row.get("source_doc_id")),
                "local_product_id": safe_text(row.get("local_product_id")),
                "local_product_name": safe_text(row.get("local_product_name")),
                "local_spec": safe_text(row.get("local_spec")),
                "local_manufacturer": safe_text(row.get("local_manufacturer")),
                "local_barcode": safe_text(row.get("local_barcode")),
                "confidence": safe_text(row.get("confidence") or candidate.get("confidence")),
                "reason_for_review": "; ".join(risk_types(row, candidate)),
                "selected_action": "create_new_product" if not product_ids else "",
                "selected_canonical_product_id": "",
                "selected_canonical_product_name": "",
                "reviewer_note": "",
                "是否确认": "否",
            }
            for index in range(3):
                number = index + 1
                item[f"candidate_{number}_product_id"] = safe_text(product_ids[index] if index < len(product_ids) else "")
                item[f"candidate_{number}_product_name"] = safe_text(product_names[index] if index < len(product_names) else "")
                item[f"candidate_{number}_spec"] = safe_text(specs[index] if index < len(specs) else "")
                item[f"candidate_{number}_evidence"] = candidate_evidence(candidate, index)
            rows.append(item)
    return rows


def load_needs_review_rows(root: Path) -> tuple[list[dict[str, str]], dict[str, dict[str, Any]]]:
    specs = [
        ("大民族", root / "config/product_identity_map.csv", root / "output/acceptance/dmz_product_mapping_candidates_full.json"),
        ("石荣霄", root / "config/product_identity_map_srx.csv", root / "output/acceptance/srx_product_mapping_candidates_full.json"),
    ]
    rows: list[dict[str, str]] = []
    fallback: dict[str, dict[str, Any]] = {}
    for brand, mapping_path, candidates_path in specs:
        candidates_by_local_id = load_candidates(candidates_path)
        for row in read_csv(mapping_path):
            if safe_text(row.get("status")) != "needs_review":
                continue
            item = {
                "brand": brand,
                "source_name": safe_text(row.get("source_name")),
                "source_doc_id": safe_text(row.get("source_doc_id")),
                "local_product_id": safe_text(row.get("local_product_id")),
                "local_product_name": safe_text(row.get("local_product_name")),
                "local_spec": safe_text(row.get("local_spec")),
                "local_manufacturer": safe_text(row.get("local_manufacturer")),
                "local_barcode": safe_text(row.get("local_barcode")),
                "confidence": safe_text(row.get("confidence")),
                "reason_for_review": safe_text(row.get("notes")) or "needs_review",
                "selected_action": "",
                "selected_canonical_product_id": "",
                "selected_canonical_product_name": "",
                "reviewer_note": "",
                "是否确认": "否",
            }
            rows.append(item)
            fallback[item["local_product_id"]] = candidates_by_local_id.get(item["local_product_id"], {})
    return rows, fallback


def source_excel_row_number(local_product_id: str) -> str:
    match = re.search(r"_r(\d+)$", safe_text(local_product_id))
    return match.group(1) if match else ""


def source_file_path(root: Path, source_name: str) -> str:
    name = safe_text(source_name)
    if not name:
        return ""
    candidates = [
        root / name,
        root / "output/product_excel_pipeline" / name,
        root / "output/product_excel_pipeline_srx" / name,
    ]
    for path in candidates:
        if path.exists():
            return str(path.relative_to(root))
    return name


def build_enhanced_review_rows(root: Path) -> list[dict[str, Any]]:
    rows, fallback = load_needs_review_rows(root)
    load_env(root)
    config = ensure_runtime_config()
    client, models_module = create_client(config)
    table_names = resolve_text_table_names(config)
    enhanced_rows: list[dict[str, Any]] = []
    for row in rows:
        candidate_payload = enhanced_candidates_for_row(
            client,
            models_module,
            table_names,
            row,
            fallback_candidate=fallback.get(row["local_product_id"]),
            top_k=10,
        )
        item: dict[str, Any] = dict(row)
        item["source_file_path"] = source_file_path(root, row.get("source_name", ""))
        item["source_excel_row_number"] = source_excel_row_number(row.get("local_product_id", ""))
        item["search_terms"] = candidate_payload["search_terms"]
        item["version_signals"] = candidate_payload["version_signals"]
        item["candidate_scores"] = candidate_payload["candidate_scores"]
        item["direct_core_hit_count"] = candidate_payload["direct_core_hit_count"]
        item["fallback_used"] = candidate_payload["fallback_used"]
        item["candidate_generation_note"] = candidate_payload["candidate_generation_note"]
        ids = candidate_payload["candidate_online_product_ids"]
        names = candidate_payload["candidate_online_product_names"]
        specs = candidate_payload["candidate_online_specs"]
        evidence_by_product = candidate_payload["candidate_evidence_by_product"]
        for index in range(10):
            number = index + 1
            product_id = ids[index] if index < len(ids) else ""
            item[f"candidate_{number}_product_id"] = product_id
            item[f"candidate_{number}_product_name"] = names[index] if index < len(names) else ""
            item[f"candidate_{number}_spec"] = specs[index] if index < len(specs) else ""
            item[f"candidate_{number}_evidence"] = "; ".join(evidence_by_product.get(product_id, [])) if product_id else ""
        enhanced_rows.append(item)
    return enhanced_rows


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=REVIEW_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "needs_review_49"
    sheet.append(REVIEW_COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in REVIEW_COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 12,
        "B": 26,
        "C": 28,
        "D": 34,
        "E": 28,
        "F": 18,
        "G": 38,
        "H": 18,
        "W": 24,
        "X": 28,
        "Y": 28,
        "Z": 34,
        "AA": 12,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column in range(9, 21):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = 28 if "evidence" not in sheet.cell(row=1, column=column).value else 36
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    action_col = REVIEW_COLUMNS.index("selected_action") + 1
    confirm_col = REVIEW_COLUMNS.index("是否确认") + 1
    action_validation = DataValidation(type="list", formula1=f'"{",".join(ACTION_OPTIONS)}"', allow_blank=True)
    confirm_validation = DataValidation(type="list", formula1=f'"{",".join(CONFIRMED_OPTIONS)}"', allow_blank=True)
    sheet.add_data_validation(action_validation)
    sheet.add_data_validation(confirm_validation)
    action_validation.add(f"{sheet.cell(row=2, column=action_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=action_col).coordinate}")
    confirm_validation.add(f"{sheet.cell(row=2, column=confirm_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=confirm_col).coordinate}")
    workbook.save(path)
    return True


def write_json(path: Path, rows: list[dict[str, str]], dropdown_added: bool) -> dict[str, Any]:
    payload = {
        "created_at": now_iso(),
        "total": len(rows),
        "brand_counts": {
            "大民族": sum(1 for row in rows if row.get("brand") == "大民族"),
            "石荣霄": sum(1 for row in rows if row.get("brand") == "石荣霄"),
        },
        "columns": REVIEW_COLUMNS,
        "selected_action_options": ACTION_OPTIONS,
        "confirmed_options": CONFIRMED_OPTIONS,
        "excel_dropdown_added": dropdown_added,
        "rows": rows,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def write_markdown(path: Path, rows: list[dict[str, str]], payload: dict[str, Any]) -> None:
    lines = [
        "# Product Mapping Review 49",
        "",
        f"- Created at: `{payload['created_at']}`",
        f"- Total needs_review: {payload['total']}",
        f"- 大民族: {payload['brand_counts']['大民族']}",
        f"- 石荣霄: {payload['brand_counts']['石荣霄']}",
        f"- Excel selected_action dropdown: `{payload['excel_dropdown_added']}`",
        "",
        "| Brand | Local product | Local spec | Candidate 1 | Candidate 2 | Candidate 3 | Confidence | Risk |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| {brand} | `{local}` | {spec} | `{c1}` {c1n} | `{c2}` {c2n} | `{c3}` {c3n} | {conf} | {risk} |".format(
                brand=row.get("brand", ""),
                local=row.get("local_product_id", ""),
                spec=row.get("local_spec", ""),
                c1=row.get("candidate_1_product_id", ""),
                c1n=row.get("candidate_1_product_name", ""),
                c2=row.get("candidate_2_product_id", ""),
                c2n=row.get("candidate_2_product_name", ""),
                c3=row.get("candidate_3_product_id", ""),
                c3n=row.get("candidate_3_product_name", ""),
                conf=row.get("confidence", ""),
                risk=row.get("reason_for_review", ""),
            )
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_summary(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        "# Product Mapping Review 49 Summary",
        "",
        f"1. 总 needs_review 数量：{payload['total']}",
        f"2. 大民族数量：{payload['brand_counts']['大民族']}",
        f"3. 石荣霄数量：{payload['brand_counts']['石荣霄']}",
        "",
        "## 需要重点看的风险类型",
        "",
        "- 同名多规格：优先核对规格、容量、套装内容。",
        "- 纪念版/年份版/生肖版：核对年份、节日、生肖和限定描述，避免映射到相近款。",
        "- 规格缺失：没有规格时不要只凭名称确认。",
        "- 候选跨品牌：候选名称不含当前品牌时需要特别谨慎。",
        "- 多个候选相近：候选 1/2/3 名称或规格接近时必须人工选择。",
        "",
        "## 如何填写 Excel",
        "",
        "1. 在 `selected_action` 中选择确认动作。",
        "2. 如果选择 `confirm_candidate_1/2/3`，把对应候选的 product_id/name 填到 `selected_canonical_product_id` 和 `selected_canonical_product_name`。",
        "3. 如果线上没有对应产品，选择 `create_new_product`。",
        "4. 资料不足选择 `need_more_info`，并在 `reviewer_note` 写明缺什么。",
        "5. 最后把 `是否确认` 改为 `是`。",
        "",
        "## 下一步如何应用回 mapping CSV",
        "",
        "填完后可按 `local_product_id` 回写到对应 `config/product_identity_map*.csv`：",
        "- `confirm_candidate_*`：更新 `canonical_product_id/name`、`status=confirmed`，保留 reviewer note 到 `notes`。",
        "- `create_new_product`：保留或生成新 canonical id，状态按后续入库策略处理。",
        "- `ignore/need_more_info`：不要生成 supplement，不改为 confirmed。",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def raw_to_zh_action(raw_value: str) -> str:
    for zh_value, value in ACTION_ZH_TO_RAW.items():
        if value == raw_value:
            return zh_value
    return ""


def zh_row(row: dict[str, str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for column in REVIEW_COLUMNS_ZH:
        if column == "我的选择":
            result[column] = raw_to_zh_action(row.get("selected_action", ""))
        elif column in ZH_COLUMN_TO_RAW:
            result[column] = row.get(ZH_COLUMN_TO_RAW[column], "")
        else:
            result[column] = ""
    return result


def write_csv_zh(path: Path, rows: list[dict[str, str]]) -> None:
    zh_rows = [zh_row(row) for row in rows]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=REVIEW_COLUMNS_ZH)
        writer.writeheader()
        writer.writerows(zh_rows)


def write_xlsx_zh(path: Path, rows: list[dict[str, str]]) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人工确认49条"
    sheet.append(REVIEW_COLUMNS_ZH)
    for row in rows:
        item = zh_row(row)
        sheet.append([item.get(column, "") for column in REVIEW_COLUMNS_ZH])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 12,
        "B": 24,
        "C": 30,
        "D": 34,
        "E": 28,
        "F": 18,
        "G": 38,
        "H": 16,
        "U": 14,
        "V": 28,
        "W": 22,
        "X": 28,
        "Y": 28,
        "Z": 12,
        "AA": 34,
        "AB": 24,
        "AC": 28,
        "AD": 28,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column in range(9, 21):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = 28 if "依据" not in sheet.cell(row=1, column=column).value else 36
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    action_col = REVIEW_COLUMNS_ZH.index("我的选择") + 1
    confirm_col = REVIEW_COLUMNS_ZH.index("是否确认") + 1
    action_validation = DataValidation(type="list", formula1=f'"{",".join(ACTION_OPTIONS_ZH)}"', allow_blank=True)
    confirm_validation = DataValidation(type="list", formula1=f'"{",".join(CONFIRMED_OPTIONS)}"', allow_blank=True)
    sheet.add_data_validation(action_validation)
    sheet.add_data_validation(confirm_validation)
    action_validation.add(f"{sheet.cell(row=2, column=action_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=action_col).coordinate}")
    confirm_validation.add(f"{sheet.cell(row=2, column=confirm_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=confirm_col).coordinate}")
    # Hide raw helper columns but keep them in the workbook for machine reading.
    for column in ["AB", "AC", "AD"]:
        sheet.column_dimensions[column].hidden = True
    workbook.save(path)
    return True


def write_json_zh(path: Path, rows: list[dict[str, str]], dropdown_added: bool) -> dict[str, Any]:
    zh_rows = [zh_row(row) for row in rows]
    payload = {
        "created_at": now_iso(),
        "total": len(rows),
        "brand_counts": {
            "大民族": sum(1 for row in rows if row.get("brand") == "大民族"),
            "石荣霄": sum(1 for row in rows if row.get("brand") == "石荣霄"),
        },
        "columns_zh": REVIEW_COLUMNS_ZH,
        "raw_columns": REVIEW_COLUMNS,
        "column_mapping_zh_to_raw": ZH_COLUMN_TO_RAW,
        "selected_action_options_zh": ACTION_OPTIONS_ZH,
        "selected_action_mapping_zh_to_raw": ACTION_ZH_TO_RAW,
        "confirmed_options": CONFIRMED_OPTIONS,
        "excel_dropdown_added": dropdown_added,
        "rows": zh_rows,
        "raw_rows": rows,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def write_markdown_zh(path: Path, rows: list[dict[str, str]], payload: dict[str, Any]) -> None:
    lines = [
        "# 产品映射人工确认表（49 条）",
        "",
        f"- 生成时间：`{payload['created_at']}`",
        f"- 总数：{payload['total']}",
        f"- 大民族：{payload['brand_counts']['大民族']}",
        f"- 石荣霄：{payload['brand_counts']['石荣霄']}",
        f"- Excel 中文下拉：`{payload['excel_dropdown_added']}`",
        "",
        "| 品牌 | 本地产品ID | 本地产品名称 | 本地规格 | 候选1 | 候选2 | 候选3 | 置信度 | 需要人工确认原因 |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| {brand} | `{local}` | {name} | {spec} | `{c1}` {c1n} | `{c2}` {c2n} | `{c3}` {c3n} | {conf} | {risk} |".format(
                brand=row.get("brand", ""),
                local=row.get("local_product_id", ""),
                name=row.get("local_product_name", ""),
                spec=row.get("local_spec", ""),
                c1=row.get("candidate_1_product_id", ""),
                c1n=row.get("candidate_1_product_name", ""),
                c2=row.get("candidate_2_product_id", ""),
                c2n=row.get("candidate_2_product_name", ""),
                c3=row.get("candidate_3_product_id", ""),
                c3n=row.get("candidate_3_product_name", ""),
                conf=row.get("confidence", ""),
                risk=row.get("reason_for_review", ""),
            )
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_chinese_instructions(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        "# 产品映射人工确认表中文说明",
        "",
        f"1. 这 49 条是当前两个映射表中 `status=needs_review` 的产品，需要人工确认本地产品是否对应某个线上 `product_id`。",
        f"2. 大民族 needs_review：{payload['brand_counts']['大民族']} 条。",
        f"3. 石荣霄 needs_review：{payload['brand_counts']['石荣霄']} 条。",
        "",
        "## 每一行应该怎么选",
        "",
        "先看本地产品名称、规格、酒厂，再比较候选 1/2/3 的产品 ID、名称、规格和匹配依据。",
        "",
        "## 什么时候选“确认候选1/2/3”",
        "",
        "- 本地产品名称和候选名称一致或高度一致。",
        "- 规格、容量、套装内容能对上。",
        "- 纪念版、年份版、生肖版、端午/中秋/周年等描述没有冲突。",
        "",
        "## 什么时候选“作为新产品”",
        "",
        "- 三个候选都明显不是同一个产品。",
        "- 线上没有对应产品，或者候选只是在品牌/关键词上相近。",
        "",
        "## 什么时候选“忽略”",
        "",
        "- 这条本地资料不需要进入线上产品库。",
        "- 重复、无效、测试、非产品记录。",
        "",
        "## 什么时候选“信息不足，稍后再看”",
        "",
        "- 名称相近但规格缺失。",
        "- 多个候选都很像，无法仅凭当前资料判断。",
        "- 需要补充图片、原 Excel、质检报告或人工业务判断。",
        "",
        "## 中文选项和程序字段对应关系",
        "",
    ]
    for zh_value, raw_value in ACTION_ZH_TO_RAW.items():
        lines.append(f"- {zh_value} -> `{raw_value}`")
    lines.extend(
        [
            "",
            "Excel 里保留了隐藏 raw 列：`selected_action_raw`、`selected_canonical_product_id`、`selected_canonical_product_name`。",
            "如果 Excel 下拉联动没有自动写 raw action，后续应用脚本会按“我的选择”映射回原始值。",
            "",
            "## 填完之后下一步怎么处理",
            "",
            "1. 把“我的选择”填好。",
            "2. 如果确认候选，填写对应线上产品 ID 和名称。",
            "3. 把“是否确认”改成“是”。",
            "4. 后续再运行专门的应用脚本，把确认结果回写到 `config/product_identity_map*.csv`。",
            "",
            "## 重要提醒",
            "",
            "没有确认的产品不会入库，也不会生成 supplement。只有确认后的映射才进入下一步处理。",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def enhanced_zh_row(row: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {
        "品牌": safe_text(row.get("brand")),
        "来源名称": safe_text(row.get("source_name")),
        "来源文件路径": safe_text(row.get("source_file_path") or row.get("source_name")),
        "源Excel行号": safe_text(row.get("source_excel_row_number")),
        "来源文件ID": safe_text(row.get("source_doc_id")),
        "本地产品ID": safe_text(row.get("local_product_id")),
        "本地产品名称": safe_text(row.get("local_product_name")),
        "本地规格": safe_text(row.get("local_spec")),
        "本地生产酒厂": safe_text(row.get("local_manufacturer")),
        "本地条码": safe_text(row.get("local_barcode")),
        "置信度": safe_text(row.get("confidence")),
        "需要人工确认原因": safe_text(row.get("reason_for_review")),
        "核心搜索词": " | ".join(row.get("search_terms") or []),
        "版本强信号": " | ".join(row.get("version_signals") or []),
        "候选生成说明": safe_text(row.get("candidate_generation_note")),
        "我的选择": "",
        "手动填写线上产品ID": "",
        "选择对应的线上产品ID": "",
        "选择对应的线上产品名称": "",
        "是否确认": "否",
        "备注": "",
        "selected_action_raw": "",
        "selected_canonical_product_id": "",
        "selected_canonical_product_name": "",
    }
    for index in range(1, 11):
        result[f"候选{index}产品ID"] = safe_text(row.get(f"candidate_{index}_product_id"))
        result[f"候选{index}产品名称"] = safe_text(row.get(f"candidate_{index}_product_name"))
        result[f"候选{index}规格"] = safe_text(row.get(f"candidate_{index}_spec"))
        result[f"候选{index}匹配依据"] = safe_text(row.get(f"candidate_{index}_evidence"))
    return result


def write_enhanced_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=REVIEW_COLUMNS_ZH_ENHANCED)
        writer.writeheader()
        writer.writerows([enhanced_zh_row(row) for row in rows])


def write_enhanced_xlsx(path: Path, rows: list[dict[str, Any]]) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "增强候选49条"
    sheet.append(REVIEW_COLUMNS_ZH_ENHANCED)
    for row in rows:
        item = enhanced_zh_row(row)
        sheet.append([item.get(column, "") for column in REVIEW_COLUMNS_ZH_ENHANCED])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index in range(1, sheet.max_column + 1):
        header = safe_text(sheet.cell(row=1, column=column_index).value)
        letter = sheet.cell(row=1, column=column_index).column_letter
        if "匹配依据" in header:
            width = 40
        elif "产品ID" in header or "来源文件ID" in header or header in {"本地产品ID"}:
            width = 30
        elif "产品名称" in header or header in {"本地生产酒厂", "核心搜索词"}:
            width = 28
        elif header in {"我的选择", "手动填写线上产品ID", "选择对应的线上产品ID", "选择对应的线上产品名称"}:
            width = 26
        else:
            width = 16
        sheet.column_dimensions[letter].width = width
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    action_col = REVIEW_COLUMNS_ZH_ENHANCED.index("我的选择") + 1
    confirm_col = REVIEW_COLUMNS_ZH_ENHANCED.index("是否确认") + 1
    action_validation = DataValidation(type="list", formula1=f'"{",".join(ACTION_OPTIONS_ZH_ENHANCED)}"', allow_blank=True)
    confirm_validation = DataValidation(type="list", formula1=f'"{",".join(CONFIRMED_OPTIONS)}"', allow_blank=True)
    sheet.add_data_validation(action_validation)
    sheet.add_data_validation(confirm_validation)
    action_validation.add(f"{sheet.cell(row=2, column=action_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=action_col).coordinate}")
    confirm_validation.add(f"{sheet.cell(row=2, column=confirm_col).coordinate}:{sheet.cell(row=max(sheet.max_row, 2), column=confirm_col).coordinate}")
    for column_index, header in enumerate(REVIEW_COLUMNS_ZH_ENHANCED, start=1):
        if header in {"selected_action_raw", "selected_canonical_product_id", "selected_canonical_product_name"}:
            sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].hidden = True
    workbook.save(path)
    return True


def write_enhanced_json(path: Path, rows: list[dict[str, Any]], dropdown_added: bool) -> dict[str, Any]:
    payload = {
        "created_at": now_iso(),
        "total": len(rows),
        "brand_counts": {
            "大民族": sum(1 for row in rows if row.get("brand") == "大民族"),
            "石荣霄": sum(1 for row in rows if row.get("brand") == "石荣霄"),
        },
        "columns_zh": REVIEW_COLUMNS_ZH_ENHANCED,
        "selected_action_options_zh": ACTION_OPTIONS_ZH_ENHANCED,
        "selected_action_mapping_zh_to_raw": ACTION_ZH_TO_RAW_ENHANCED,
        "confirmed_options": CONFIRMED_OPTIONS,
        "excel_dropdown_added": dropdown_added,
        "candidate_limit": 10,
        "search_strategy": "core_product_name_first_readonly_text_docs",
        "rows": [enhanced_zh_row(row) for row in rows],
        "raw_rows": rows,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


SAMPLE_NAMES = [
    "大民族酒·师藏年鉴（陈任远版）",
    "大民族酒·师藏年鉴（张怀仁版）",
    "大民族酒·天字号V30",
    "大民族酒·天字号A30",
    "大民族酒·蛇年生肖纪念酒",
    "大民族酒·端午纪念",
]


def write_enhanced_markdown(path: Path, rows: list[dict[str, Any]], payload: dict[str, Any]) -> None:
    lines = [
        "# 产品映射人工确认表（增强版）",
        "",
        f"- 生成时间：`{payload['created_at']}`",
        f"- 总数：{payload['total']}",
        f"- 大民族：{payload['brand_counts']['大民族']}",
        f"- 石荣霄：{payload['brand_counts']['石荣霄']}",
        "- 搜索策略：优先按产品核心名称只读查询线上 `text_docs`。",
        "",
        "| 品牌 | 本地产品 | 核心搜索词 | 候选1 | 候选2 | 候选3 | 候选4 | 候选5 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| {brand} | {name} | {terms} | {c1} | {c2} | {c3} | {c4} | {c5} |".format(
                brand=row.get("brand", ""),
                name=row.get("local_product_name", ""),
                terms=" / ".join((row.get("search_terms") or [])[:3]),
                c1=f"`{row.get('candidate_1_product_id','')}` {row.get('candidate_1_product_name','')}",
                c2=f"`{row.get('candidate_2_product_id','')}` {row.get('candidate_2_product_name','')}",
                c3=f"`{row.get('candidate_3_product_id','')}` {row.get('candidate_3_product_name','')}",
                c4=f"`{row.get('candidate_4_product_id','')}` {row.get('candidate_4_product_name','')}",
                c5=f"`{row.get('candidate_5_product_id','')}` {row.get('candidate_5_product_name','')}",
            )
        )
    lines.extend(["", "## 重点样例候选 1-5", ""])
    for sample in SAMPLE_NAMES:
        matched = [row for row in rows if row.get("local_product_name") == sample]
        if not matched:
            continue
        row = matched[0]
        lines.extend([f"### {sample}", ""])
        for index in range(1, 6):
            lines.append(
                f"- 候选{index}: `{row.get(f'candidate_{index}_product_id','')}` "
                f"{row.get(f'candidate_{index}_product_name','')} / {row.get(f'candidate_{index}_spec','')} "
                f"({row.get(f'candidate_{index}_evidence','')})"
            )
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_enhanced_instructions(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        "# 产品映射增强版确认表说明",
        "",
        "1. 现在候选优先按产品核心名称搜索，不再优先用“纪念版 / 年份版 / 规格 / 生肖”等泛词拆开搜索。",
        "2. 搜索会保留括号/版本强信号，例如陈任远版、张怀仁版、V30、A15、精品10号等。",
        "3. 如果候选里还是没有正确产品，可以在“我的选择”里选“手动填写线上产品ID”，再填入线上产品 ID。",
        "4. 如果线上确实没有，就选“作为新产品”。",
        "5. 如果无法判断，就选“信息不足，稍后再看”。",
        "",
        "## 中文选项和程序字段对应关系",
        "",
    ]
    for zh_value, raw_value in ACTION_ZH_TO_RAW_ENHANCED.items():
        lines.append(f"- {zh_value} -> `{raw_value}`")
    lines.extend(
        [
            "",
            "本轮没有自动确认任何产品，没有生成 supplement，也没有写 OpenSearch。",
            "",
            "## 填写建议",
            "",
            "- 优先看候选 1-3；若都不对，再看候选 4-10。",
            "- 如果你知道线上 product_id 但候选没列出，选择“手动填写线上产品ID”。",
            "- 没有确认的产品不会入库。",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_enhanced_outputs(root: Path, output_dir: Path) -> dict[str, Any]:
    rows = build_enhanced_review_rows(root)
    csv_path = output_dir / "product_mapping_review_49_增强版.csv"
    xlsx_path = output_dir / "product_mapping_review_49_增强版.xlsx"
    json_path = output_dir / "product_mapping_review_49_增强版.json"
    md_path = output_dir / "product_mapping_review_49_增强版.md"
    instructions_path = output_dir / "product_mapping_review_49_增强版说明.md"
    write_enhanced_csv(csv_path, rows)
    dropdown_added = write_enhanced_xlsx(xlsx_path, rows)
    payload = write_enhanced_json(json_path, rows, dropdown_added)
    write_enhanced_markdown(md_path, rows, payload)
    write_enhanced_instructions(instructions_path, payload)
    return {
        "rows": rows,
        "payload": payload,
        "paths": {
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "json": str(json_path),
            "md": str(md_path),
            "instructions": str(instructions_path),
        },
    }


def build_chinese_outputs(root: Path, output_dir: Path) -> dict[str, Any]:
    rows = build_review_rows(
        [
            (
                "大民族",
                root / "config/product_identity_map.csv",
                root / "output/acceptance/dmz_product_mapping_candidates_full.json",
            ),
            (
                "石荣霄",
                root / "config/product_identity_map_srx.csv",
                root / "output/acceptance/srx_product_mapping_candidates_full.json",
            ),
        ]
    )
    csv_path = output_dir / "product_mapping_review_49_中文.csv"
    xlsx_path = output_dir / "product_mapping_review_49_中文.xlsx"
    json_path = output_dir / "product_mapping_review_49_中文.json"
    md_path = output_dir / "product_mapping_review_49_中文.md"
    instructions_path = output_dir / "product_mapping_review_49_中文说明.md"
    write_csv_zh(csv_path, rows)
    dropdown_added = write_xlsx_zh(xlsx_path, rows)
    payload = write_json_zh(json_path, rows, dropdown_added)
    write_markdown_zh(md_path, rows, payload)
    write_chinese_instructions(instructions_path, payload)
    return {
        "rows": rows,
        "payload": payload,
        "paths": {
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "json": str(json_path),
            "md": str(md_path),
            "instructions": str(instructions_path),
        },
    }


def build_outputs(root: Path, output_dir: Path) -> dict[str, Any]:
    rows = build_review_rows(
        [
            (
                "大民族",
                root / "config/product_identity_map.csv",
                root / "output/acceptance/dmz_product_mapping_candidates_full.json",
            ),
            (
                "石荣霄",
                root / "config/product_identity_map_srx.csv",
                root / "output/acceptance/srx_product_mapping_candidates_full.json",
            ),
        ]
    )
    csv_path = output_dir / "product_mapping_review_49.csv"
    xlsx_path = output_dir / "product_mapping_review_49.xlsx"
    json_path = output_dir / "product_mapping_review_49.json"
    md_path = output_dir / "product_mapping_review_49.md"
    summary_path = output_dir / "product_mapping_review_49_summary.md"
    write_csv(csv_path, rows)
    dropdown_added = write_xlsx(xlsx_path, rows)
    payload = write_json(json_path, rows, dropdown_added)
    write_markdown(md_path, rows, payload)
    write_summary(summary_path, payload)
    return {
        "rows": rows,
        "payload": payload,
        "paths": {
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "json": str(json_path),
            "md": str(md_path),
            "summary": str(summary_path),
        },
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build manual product mapping review files.")
    parser.add_argument("--root", default=str(ROOT))
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--chinese", action="store_true", help="Build Chinese display copies with raw machine-readable fields.")
    parser.add_argument("--enhanced", action="store_true", help="Build enhanced needs_review table using core product-name search.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = Path(args.root).resolve()
    output_dir = (root / args.output_dir).resolve() if not Path(args.output_dir).is_absolute() else Path(args.output_dir).resolve()
    if args.enhanced:
        result = build_enhanced_outputs(root, output_dir)
    elif args.chinese:
        result = build_chinese_outputs(root, output_dir)
    else:
        result = build_outputs(root, output_dir)
    payload = result["payload"]
    print(
        json.dumps(
            {
                "total": payload["total"],
                "brand_counts": payload["brand_counts"],
                "excel_dropdown_added": payload["excel_dropdown_added"],
                "paths": result["paths"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
