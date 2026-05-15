#!/usr/bin/env python3
"""
Clean and standardize Excel workbooks into structured product records.

This script scans the current project for Excel files, expands merged cells,
removes obviously empty rows/columns, standardizes common field names, keeps
source traceability, and exports JSON / CSV / Markdown report artifacts.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    import pandas as pd  # type: ignore
except Exception:  # pandas is only needed for .xls probing and XLSX export.
    pd = None


ROOT = Path(__file__).resolve().parents[1]
SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".xlsm"}
PLACEHOLDER_VALUES = {
    "",
    "-",
    "--",
    "---",
    "—",
    "——",
    "———",
    "暂无",
    "无",
    "未填写",
    "未标注",
}
IMAGE_MARKER = "embedded_image"

PRIORITY_FIELDS = [
    "product_name",
    "brand",
    "series",
    "spec",
    "grade",
    "channel",
    "price",
    "packaging_desc",
    "quality_desc",
    "selling_points",
    "product_story",
    "target_users",
    "gift_attributes",
]

ADDITIONAL_FIELDS = [
    "manufacturer",
    "aroma_type",
    "positioning",
    "benchmark_product",
    "image_reference",
]

FIELD_ORDER = [
    "source_file",
    "source_sheet",
    "source_row",
    *PRIORITY_FIELDS,
    *ADDITIONAL_FIELDS,
    "has_image",
    "image_anchor_count",
    "image_columns",
    "extra_fields",
]


HEADER_ALIASES = {
    "product_name": {
        "产品名称",
        "品名",
        "产品名",
    },
    "brand": {
        "品牌",
        "品牌名称",
    },
    "series": {
        "系列",
        "产品系列",
    },
    "spec": {
        "规格",
        "产品规格",
        "容量",
        "产品容量",
    },
    "grade": {
        "产品档次",
        "档次",
        "等级",
        "产品档次中高低",
    },
    "channel": {
        "售卖渠道",
        "销售渠道",
        "产品售卖渠道",
        "渠道",
    },
    "price": {
        "价格",
        "产品价格",
        "扫码价格定位价格实际活动售卖价",
        "定位价格",
        "活动售卖价",
        "售价",
    },
    "packaging_desc": {
        "包装",
        "包装描述",
        "包装特点",
        "包装特点工艺设计理念实物颜色",
        "包装特点工艺",
        "包装工艺",
    },
    "quality_desc": {
        "产品酒质介绍",
        "酒质介绍",
        "酒体介绍",
        "酒体",
        "酒质",
        "品质描述",
        "口感描述",
    },
    "selling_points": {
        "设计理念卖点",
        "设计理念产品卖点",
        "产品卖点",
        "卖点",
        "设计理念",
        "设计理念卖点包装",
    },
    "product_story": {
        "产品故事",
        "品牌故事",
        "产品名寓意",
        "寓意",
    },
    "target_users": {
        "产品面向的销售人群及特点",
        "面向的销售人群及特点",
        "目标人群",
        "适用人群",
        "销售人群",
        "消费人群",
    },
    "gift_attributes": {
        "礼赠属性",
        "送礼属性",
        "赠礼属性",
        "礼品属性",
    },
    "manufacturer": {
        "生产酒厂",
        "生产灌装厂家",
        "生产厂家",
        "厂家",
        "酒厂",
    },
    "aroma_type": {
        "香型",
    },
    "positioning": {
        "产品定位",
        "定位",
        "产品定位引流品私域承接品",
    },
    "benchmark_product": {
        "产品对标竞品",
        "对标竞品",
        "竞品",
    },
    "image_reference": {
        "图片",
        "产品图片",
        "产品多角度图片",
        "产品图片设计图效果图",
        "产品多角度图片设计图效果图",
    },
}

HEADER_KEYWORDS = (
    "产品",
    "包装",
    "规格",
    "档次",
    "定位",
    "图片",
    "渠道",
    "酒质",
    "酒厂",
    "寓意",
    "竞品",
    "卖点",
    "设计理念",
    "故事",
    "香型",
)
TITLE_ROW_KEYWORDS = ("基础资料", "编写前", "由产品部提供")


def normalize_header_text(value: str) -> str:
    """Normalize a header cell so loosely formatted synonyms can still match."""
    value = str(value or "")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[\s\n\t\xa0]+", "", value)
    value = re.sub(r"[，,：:；;（()）【】\[\]/\\|·•、\-—_]+", "", value)
    return value


NORMALIZED_ALIAS_TO_FIELD = {
    normalize_header_text(alias): field
    for field, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


def normalize_linebreaks(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_scalar(value: Any) -> Any:
    """Lightweight cleaning that keeps business content but removes placeholders."""
    if value is None:
        return None
    if isinstance(value, str):
        if is_image_formula(value):
            return None
        cleaned = normalize_linebreaks(value)
        if cleaned in PLACEHOLDER_VALUES:
            return None
        return cleaned or None
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if value.is_integer():
            return int(value)
    return value


def is_image_formula(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    return "DISPIMG(" in value.upper()


def has_meaningful_value(value: Any) -> bool:
    return value is not None and value != ""


def discover_excel_files(root: Path) -> list[Path]:
    files = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            files.append(path)
    return sorted(files)


def build_worksheet_matrix(ws) -> list[list[Any]]:
    """Read sheet cells into memory and backfill merged ranges."""
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    matrix = [[ws.cell(row=r, column=c).value for c in range(1, max_col + 1)] for r in range(1, max_row + 1)]

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = matrix[min_row - 1][min_col - 1]
        for row_idx in range(min_row - 1, max_row):
            for col_idx in range(min_col - 1, max_col):
                matrix[row_idx][col_idx] = top_left_value

    return matrix


def extract_image_anchors(ws) -> dict[int, set[int]]:
    """Return row -> columns for embedded images, based on their anchors."""
    row_to_cols: dict[int, set[int]] = defaultdict(set)
    for image in getattr(ws, "_images", []):
        anchor = getattr(image, "anchor", None)
        if anchor is None:
            continue
        try:
            row = anchor._from.row + 1
            col = anchor._from.col + 1
        except Exception:
            continue
        row_to_cols[row].add(col)
    return row_to_cols


def map_column_name(raw_header: Any) -> str | None:
    if raw_header is None:
        return None
    raw_text = clean_scalar(raw_header)
    if not isinstance(raw_text, str):
        return None
    normalized = normalize_header_text(raw_text)
    if not normalized:
        return None

    direct = NORMALIZED_ALIAS_TO_FIELD.get(normalized)
    if direct:
        return direct

    if any(keyword in normalized for keyword in TITLE_ROW_KEYWORDS):
        return None

    # Long narrative title rows are common in these workbooks. Treat them as
    # titles, not headers, unless they hit an explicit alias above.
    if len(normalized) > 25:
        return None

    fallback_checks = [
        ("product_name", ("产品名称", "品名")),
        ("brand", ("品牌",)),
        ("series", ("系列",)),
        ("spec", ("规格", "容量")),
        ("grade", ("档次", "等级")),
        ("channel", ("售卖渠道", "销售渠道", "渠道")),
        ("price", ("价格", "售价")),
        ("packaging_desc", ("包装",)),
        ("quality_desc", ("酒质", "酒体", "口感")),
        ("selling_points", ("卖点", "设计理念")),
        ("product_story", ("故事", "寓意")),
        ("target_users", ("人群", "用户")),
        ("gift_attributes", ("礼赠", "送礼")),
        ("manufacturer", ("酒厂", "厂家")),
        ("aroma_type", ("香型",)),
        ("positioning", ("定位",)),
        ("benchmark_product", ("竞品",)),
        ("image_reference", ("图片", "效果图", "设计图")),
    ]
    for field_name, keywords in fallback_checks:
        if all(keyword in normalized for keyword in keywords[:1]) and any(keyword in normalized for keyword in keywords):
            return field_name
    return None


def score_header_row(row_values: list[Any]) -> tuple[int, int]:
    nonempty_values = [value for value in row_values if has_meaningful_value(value)]
    if not nonempty_values:
        return 0, 0

    unique_texts = {
        str(value).strip()
        for value in nonempty_values
        if isinstance(value, str) and str(value).strip()
    }
    mapped_count = 0
    score = min(len(unique_texts) or len(nonempty_values), 12)
    for value in nonempty_values:
        if not isinstance(value, str):
            continue
        mapped = map_column_name(value)
        text = value.strip()
        if mapped:
            mapped_count += 1
            score += 8
        elif len(text) <= 20 and any(keyword in text for keyword in HEADER_KEYWORDS):
            score += 2
        if len(text) >= 50:
            score -= 2
    if len(unique_texts) <= 2 and len(nonempty_values) >= 4:
        score -= 20
    if len(unique_texts) == 1 and len(nonempty_values) >= 6:
        return -50, 0
    return score, mapped_count


def detect_header_row(clean_matrix: list[list[Any]], max_scan_rows: int = 10) -> tuple[int | None, int, int]:
    best_row_index: int | None = None
    best_score = -1
    best_mapped = 0
    for row_index, row_values in enumerate(clean_matrix[:max_scan_rows]):
        score, mapped_count = score_header_row(row_values)
        if score > best_score or (score == best_score and mapped_count > best_mapped):
            best_row_index = row_index
            best_score = score
            best_mapped = mapped_count
    if best_row_index is None or best_score < 10 or best_mapped < 2:
        return None, best_score, best_mapped
    return best_row_index, best_score, best_mapped


def compute_effective_bounds(
    clean_matrix: list[list[Any]],
    formula_image_cells: set[tuple[int, int]],
    image_anchor_map: dict[int, set[int]],
) -> tuple[int, int, int, int] | None:
    min_row = min_col = None
    max_row = max_col = None

    def touch(row_number: int, col_number: int) -> None:
        nonlocal min_row, min_col, max_row, max_col
        if min_row is None or row_number < min_row:
            min_row = row_number
        if max_row is None or row_number > max_row:
            max_row = row_number
        if min_col is None or col_number < min_col:
            min_col = col_number
        if max_col is None or col_number > max_col:
            max_col = col_number

    for row_number, row_values in enumerate(clean_matrix, start=1):
        for col_number, value in enumerate(row_values, start=1):
            if has_meaningful_value(value):
                touch(row_number, col_number)

    for row_number, col_number in formula_image_cells:
        touch(row_number, col_number)

    for row_number, columns in image_anchor_map.items():
        for col_number in columns:
            touch(row_number, col_number)

    if min_row is None or min_col is None or max_row is None or max_col is None:
        return None
    return min_row, min_col, max_row, max_col


def bounds_to_range(bounds: tuple[int, int, int, int] | None) -> str:
    if bounds is None:
        return "N/A"
    min_row, min_col, max_row, max_col = bounds
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def coalesce_text(existing: Any, new_value: Any) -> Any:
    if not has_meaningful_value(existing):
        return new_value
    if not has_meaningful_value(new_value):
        return existing
    if existing == new_value:
        return existing
    if isinstance(existing, str) and isinstance(new_value, str):
        existing_parts = [part.strip() for part in existing.split("\n") if part.strip()]
        if new_value not in existing_parts:
            return existing + "\n" + new_value
    return existing


def infer_brand_and_series(record: dict[str, Any]) -> None:
    product_name = record.get("product_name")
    if not isinstance(product_name, str) or not product_name:
        return

    if not record.get("brand"):
        if "·" in product_name:
            candidate = product_name.split("·", 1)[0].strip("（( ").strip()
            if 1 <= len(candidate) <= 20:
                record["brand"] = candidate
        elif "（" in product_name:
            candidate = product_name.split("（", 1)[0].strip()
            if 1 <= len(candidate) <= 20:
                record["brand"] = candidate

    if not record.get("series") and "·" in product_name:
        candidate = product_name.split("·", 1)[1].strip("）) ").strip()
        if candidate:
            record["series"] = candidate


def infer_gift_attributes(record: dict[str, Any]) -> None:
    if record.get("gift_attributes"):
        return
    corpus = "\n".join(
        str(value)
        for value in (
            record.get("target_users"),
            record.get("selling_points"),
            record.get("packaging_desc"),
            record.get("product_story"),
        )
        if value
    )
    keywords = []
    for keyword in ("礼赠", "送礼", "收藏", "品鉴", "宴请", "聚会", "非卖品", "纪念", "节庆"):
        if keyword in corpus:
            keywords.append(keyword)
    if keywords:
        record["gift_attributes"] = "、".join(dict.fromkeys(keywords))


def normalize_record(record: dict[str, Any]) -> dict[str, Any]:
    infer_brand_and_series(record)
    infer_gift_attributes(record)
    return record


def make_sheet_issue(
    relative_file: str,
    sheet_name: str,
    reason: str,
    status: str = "abnormal",
) -> dict[str, Any]:
    return {
        "file": relative_file,
        "sheet": sheet_name,
        "status": status,
        "reason": reason,
    }


def analyze_sheet(ws, relative_file: str, global_column_mapping: dict[str, str]) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    raw_matrix = build_worksheet_matrix(ws)
    formula_image_cells: set[tuple[int, int]] = set()
    clean_matrix: list[list[Any]] = []
    for row_number, row_values in enumerate(raw_matrix, start=1):
        cleaned_row = []
        for col_number, value in enumerate(row_values, start=1):
            if is_image_formula(value):
                formula_image_cells.add((row_number, col_number))
                cleaned_row.append(None)
            else:
                cleaned_row.append(clean_scalar(value))
        clean_matrix.append(cleaned_row)

    image_anchor_map = extract_image_anchors(ws)
    effective_bounds = compute_effective_bounds(clean_matrix, formula_image_cells, image_anchor_map)
    effective_range = bounds_to_range(effective_bounds)

    header_row_index, header_score, header_mapped_count = detect_header_row(clean_matrix)
    sheet_scan = {
        "sheet_name": ws.title,
        "effective_range": effective_range,
        "merged_cells": len(ws.merged_cells.ranges),
        "embedded_images": len(getattr(ws, "_images", [])),
        "header_row": header_row_index + 1 if header_row_index is not None else None,
        "header_score": header_score,
        "recognized_header_fields": [],
        "record_count": 0,
        "status": "ok",
        "notes": [],
    }

    issues: list[dict[str, Any]] = []
    if header_row_index is None:
        sheet_scan["status"] = "abnormal"
        reason = "未能可靠识别表头行，已跳过解析。"
        sheet_scan["notes"].append(reason)
        issues.append(make_sheet_issue(relative_file, ws.title, reason))
        return [], sheet_scan, issues

    header_row = clean_matrix[header_row_index]
    active_columns: list[int] = []
    recognized_fields: list[str] = []

    for col_number, header_value in enumerate(header_row, start=1):
        column_has_text_data = any(
            has_meaningful_value(clean_matrix[row_index][col_number - 1])
            for row_index in range(header_row_index + 1, len(clean_matrix))
        )
        if not has_meaningful_value(header_value) and not column_has_text_data:
            continue
        active_columns.append(col_number)
        mapped_field = map_column_name(header_value)
        if isinstance(header_value, str):
            global_column_mapping.setdefault(header_value, mapped_field or "extra_fields")
        if mapped_field:
            recognized_fields.append(mapped_field)

    sheet_scan["recognized_header_fields"] = sorted(dict.fromkeys(recognized_fields))

    if "product_name" not in sheet_scan["recognized_header_fields"]:
        reason = "表头中未识别到产品名称列，已跳过解析。"
        sheet_scan["status"] = "abnormal"
        sheet_scan["notes"].append(reason)
        issues.append(make_sheet_issue(relative_file, ws.title, reason))
        return [], sheet_scan, issues

    image_header_columns = {
        col_number
        for col_number in active_columns
        if map_column_name(header_row[col_number - 1]) == "image_reference"
    }

    records: list[dict[str, Any]] = []
    business_value_counts: list[int] = []
    for row_number in range(header_row_index + 2, len(clean_matrix) + 1):
        row_values = clean_matrix[row_number - 1]
        anchor_cols = image_anchor_map.get(row_number, set())
        row_image_cols = set(anchor_cols)
        row_image_cols.update(
            col_number for row_idx, col_number in formula_image_cells if row_idx == row_number
        )

        is_empty_row = True
        for col_number in active_columns:
            if has_meaningful_value(row_values[col_number - 1]):
                is_empty_row = False
                break
        if is_empty_row and not row_image_cols:
            continue

        record = {field: None for field in PRIORITY_FIELDS + ADDITIONAL_FIELDS}
        extra_fields: dict[str, Any] = {}
        filled_business_fields = 0

        for col_number in active_columns:
            raw_header = header_row[col_number - 1]
            header_name = (
                raw_header
                if isinstance(raw_header, str) and raw_header
                else f"unnamed_col_{get_column_letter(col_number)}"
            )
            field_name = map_column_name(raw_header)
            cell_value = row_values[col_number - 1]

            if has_meaningful_value(cell_value):
                if field_name:
                    record[field_name] = coalesce_text(record.get(field_name), cell_value)
                    if field_name not in {"image_reference"}:
                        filled_business_fields += 1
                else:
                    extra_fields[header_name] = cell_value
                    if normalize_header_text(header_name) != "序号":
                        filled_business_fields += 1

            if col_number in image_header_columns and (has_meaningful_value(cell_value) or col_number in row_image_cols):
                record["image_reference"] = coalesce_text(record.get("image_reference"), cell_value or IMAGE_MARKER)

        if row_image_cols and not record.get("image_reference"):
            record["image_reference"] = IMAGE_MARKER

        product_name = record.get("product_name")
        if not has_meaningful_value(product_name):
            continue

        normalized_record = normalize_record(
            {
                "source_file": relative_file,
                "source_sheet": ws.title,
                "source_row": row_number,
                **record,
                "has_image": bool(row_image_cols or record.get("image_reference")),
                "image_anchor_count": len(anchor_cols),
                "image_columns": [get_column_letter(col) for col in sorted(row_image_cols)] if row_image_cols else [],
                "extra_fields": extra_fields,
            }
        )
        records.append(normalized_record)
        business_value_counts.append(filled_business_fields)

    sheet_scan["record_count"] = len(records)

    if sheet_scan["header_row"] and sheet_scan["header_row"] > 3:
        sheet_scan["notes"].append("表头位于第 4 行以后，结构偏特殊。")

    if len(records) == 0:
        reason = "识别到表头，但没有抽取到产品记录。"
        sheet_scan["status"] = "abnormal"
        sheet_scan["notes"].append(reason)
        issues.append(make_sheet_issue(relative_file, ws.title, reason))
        return records, sheet_scan, issues

    average_business_fields = sum(business_value_counts) / len(business_value_counts) if business_value_counts else 0.0
    if average_business_fields < 3 or len(sheet_scan["recognized_header_fields"]) < 4:
        reason = (
            f"结构较稀疏，平均每条记录仅识别 {average_business_fields:.1f} 个业务字段，"
            f"建议人工复核。"
        )
        sheet_scan["status"] = "review"
        sheet_scan["notes"].append(reason)
        issues.append(make_sheet_issue(relative_file, ws.title, reason, status="review"))

    return records, sheet_scan, issues


def process_xlsx_workbook(
    workbook_path: Path,
    root: Path,
    global_column_mapping: dict[str, str],
) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    relative_file = str(workbook_path.relative_to(root))
    workbook_scan = {
        "file": relative_file,
        "sheet_count": 0,
        "sheets": [],
        "status": "ok",
    }
    records: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []

    workbook = load_workbook(workbook_path, data_only=False)
    workbook_scan["sheet_count"] = len(workbook.worksheets)
    try:
        for ws in workbook.worksheets:
            sheet_records, sheet_scan, sheet_issues = analyze_sheet(ws, relative_file, global_column_mapping)
            records.extend(sheet_records)
            workbook_scan["sheets"].append(sheet_scan)
            issues.extend(sheet_issues)
            if sheet_scan["status"] == "abnormal":
                workbook_scan["status"] = "abnormal"
        return records, workbook_scan, issues
    finally:
        workbook.close()


def process_xls_workbook(
    workbook_path: Path,
    root: Path,
    global_column_mapping: dict[str, str],
) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    relative_file = str(workbook_path.relative_to(root))
    workbook_scan = {
        "file": relative_file,
        "sheet_count": 0,
        "sheets": [],
        "status": "abnormal",
    }
    issues: list[dict[str, Any]] = []

    try:
        if pd is None:
            raise RuntimeError("未安装 pandas")
        excel_file = pd.ExcelFile(workbook_path)
    except Exception as exc:
        reason = f".xls 文件当前环境无法解析：{exc}"
        issues.append(make_sheet_issue(relative_file, "(workbook)", reason))
        workbook_scan["sheets"].append(
            {
                "sheet_name": "(workbook)",
                "effective_range": "N/A",
                "merged_cells": None,
                "embedded_images": None,
                "header_row": None,
                "header_score": None,
                "recognized_header_fields": [],
                "record_count": 0,
                "status": "abnormal",
                "notes": [reason],
            }
        )
        return [], workbook_scan, issues

    workbook_scan["sheet_count"] = len(excel_file.sheet_names)
    for sheet_name in excel_file.sheet_names:
        global_column_mapping.setdefault(f"{sheet_name}(.xls)", "extra_fields")
        reason = ".xls 解析能力受限，当前脚本未对该 sheet 做结构化抽取。"
        workbook_scan["sheets"].append(
            {
                "sheet_name": sheet_name,
                "effective_range": "N/A",
                "merged_cells": None,
                "embedded_images": None,
                "header_row": None,
                "header_score": None,
                "recognized_header_fields": [],
                "record_count": 0,
                "status": "abnormal",
                "notes": [reason],
            }
        )
        issues.append(make_sheet_issue(relative_file, sheet_name, reason))
    return [], workbook_scan, issues


def flatten_record_for_csv(record: dict[str, Any]) -> dict[str, Any]:
    flat = {field: record.get(field) for field in FIELD_ORDER if field != "extra_fields"}
    flat["image_columns"] = ",".join(record.get("image_columns") or [])
    flat["extra_fields"] = json.dumps(record.get("extra_fields") or {}, ensure_ascii=False, sort_keys=True)
    return flat


def write_csv_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx_rows_if_possible(path: Path, rows: list[dict[str, Any]]) -> None:
    if pd is not None:
        pd.DataFrame(rows).to_excel(path, index=False)
        return
    try:
        from openpyxl import Workbook

        workbook = Workbook()
        ws = workbook.active
        ws.title = "products_cleaned"
        fieldnames: list[str] = []
        seen: set[str] = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    fieldnames.append(key)
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(key) for key in fieldnames])
        workbook.save(path)
    except Exception:
        # JSON/CSV are the authoritative downstream artifacts for this pipeline.
        return


def normalize_duplicate_key_part(value: Any) -> str:
    text = normalize_header_text("" if value is None else str(value))
    return text.lower()


def business_field_count(record: dict[str, Any]) -> int:
    fields = PRIORITY_FIELDS + ADDITIONAL_FIELDS
    return sum(1 for field_name in fields if has_meaningful_value(record.get(field_name)))


def build_products_check_report(records: list[dict[str, Any]], issues: list[dict[str, Any]]) -> str:
    field_counts = Counter()
    for record in records:
        for field_name in PRIORITY_FIELDS + ADDITIONAL_FIELDS:
            if has_meaningful_value(record.get(field_name)):
                field_counts[field_name] += 1

    key_missing_rows: list[str] = []
    sparse_rows: list[str] = []
    extra_field_rows: list[str] = []
    duplicate_counter: Counter[tuple[str, str, str]] = Counter()
    duplicate_examples: dict[tuple[str, str, str], list[str]] = defaultdict(list)

    for record in records:
        source = f"{record['source_file']} / {record['source_sheet']} / row {record['source_row']}"
        missing_fields = [
            field_name
            for field_name in ["product_name", "brand", "spec", "packaging_desc", "selling_points"]
            if not has_meaningful_value(record.get(field_name))
        ]
        if missing_fields:
            key_missing_rows.append(
                f"- `{record.get('product_name') or '(无产品名)'}`：缺少 {', '.join(missing_fields)}；来源：{source}"
            )

        populated = business_field_count(record)
        if populated <= 5:
            sparse_rows.append(
                f"- `{record.get('product_name') or '(无产品名)'}`：仅识别 {populated} 个业务字段；来源：{source}"
            )

        extra_fields = record.get("extra_fields") or {}
        if isinstance(extra_fields, dict) and extra_fields:
            extra_field_rows.append(
                f"- `{record.get('product_name') or '(无产品名)'}`：异常/未归一字段 {', '.join(extra_fields.keys())}；来源：{source}"
            )

        duplicate_key = (
            normalize_duplicate_key_part(record.get("product_name")),
            normalize_duplicate_key_part(record.get("brand")),
            normalize_duplicate_key_part(record.get("spec")),
        )
        if duplicate_key[0]:
            duplicate_counter[duplicate_key] += 1
            if len(duplicate_examples[duplicate_key]) < 5:
                duplicate_examples[duplicate_key].append(source)

    duplicate_lines = []
    for duplicate_key, count in duplicate_counter.items():
        if count < 2:
            continue
        product_name, brand, spec = duplicate_key
        duplicate_lines.append(
            f"- `{product_name or '(空)'}` / `{brand or '(空)'}` / `{spec or '(空)'}`：{count} 条；"
            f"样例来源：{'；'.join(duplicate_examples[duplicate_key])}"
        )

    if not duplicate_lines:
        duplicate_lines.append("- 未发现基于 `product_name + brand + spec` 的重复产品。")
    if not key_missing_rows:
        key_missing_rows.append("- 无。")
    if not sparse_rows:
        sparse_rows.append("- 无。")
    if not extra_field_rows:
        extra_field_rows.append("- 无。")

    issue_lines = [
        f"- `{issue['status']}` {issue['file']} / {issue['sheet']}: {issue['reason']}"
        for issue in issues
    ]
    if not issue_lines:
        issue_lines.append("- 无。")

    fill_rate_lines = ["| 字段 | 非空数 | 填充率 |", "| --- | ---: | ---: |"]
    total = len(records) or 1
    for field_name in PRIORITY_FIELDS + ADDITIONAL_FIELDS:
        count = field_counts.get(field_name, 0)
        fill_rate_lines.append(f"| `{field_name}` | {count} | {count / total:.1%} |")

    return f"""# 产品提取检查报告

## 总览

- 产品总数：{len(records)}
- 结构异常/需复核 sheet 数：{len(issues)}
- 关键字段缺失记录数：{len(key_missing_rows) if key_missing_rows != ['- 无。'] else 0}
- 稀疏记录数（业务字段 <= 5）：{len(sparse_rows) if sparse_rows != ['- 无。'] else 0}
- 含异常/未归一字段记录数：{len(extra_field_rows) if extra_field_rows != ['- 无。'] else 0}
- 重复产品组数：{sum(1 for count in duplicate_counter.values() if count >= 2)}

## 字段填充率

{chr(10).join(fill_rate_lines)}

## 结构异常或需复核 Sheet

{chr(10).join(issue_lines)}

## 关键字段缺失记录

{chr(10).join(key_missing_rows[:50])}

## 稀疏记录

{chr(10).join(sparse_rows[:50])}

## 异常/未归一字段

{chr(10).join(extra_field_rows[:50])}

## 重复产品检查

说明：按 `product_name + brand + spec` 归一后判重，仅用于人工排查。

{chr(10).join(duplicate_lines[:50])}
"""


def render_sample_table(records: list[dict[str, Any]], limit: int = 5) -> str:
    sample_records = records[:limit]
    if not sample_records:
        return "无样本可展示。"

    lines = [
        "| product_name | brand | spec | channel | source |",
        "| --- | --- | --- | --- | --- |",
    ]
    for record in sample_records:
        source = f"{record['source_file']} / {record['source_sheet']} / row {record['source_row']}"
        lines.append(
            "| {product_name} | {brand} | {spec} | {channel} | {source} |".format(
                product_name=str(record.get("product_name") or "").replace("\n", " / "),
                brand=str(record.get("brand") or "").replace("\n", " / "),
                spec=str(record.get("spec") or "").replace("\n", " / "),
                channel=str(record.get("channel") or "").replace("\n", " / "),
                source=source.replace("|", "/"),
            )
        )
    return "\n".join(lines)


def render_scan_table(workbook_scans: list[dict[str, Any]]) -> str:
    lines = [
        "| 文件 | Sheet | 有效区域 | 表头行 | 记录数 | 合并单元格 | 嵌入图片 | 状态 | 备注 |",
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for workbook in workbook_scans:
        for sheet in workbook["sheets"]:
            lines.append(
                "| {file} | {sheet_name} | {effective_range} | {header_row} | {record_count} | {merged_cells} | {embedded_images} | {status} | {notes} |".format(
                    file=workbook["file"],
                    sheet_name=sheet["sheet_name"],
                    effective_range=sheet["effective_range"],
                    header_row=sheet["header_row"] if sheet["header_row"] is not None else "",
                    record_count=sheet["record_count"],
                    merged_cells=sheet["merged_cells"] if sheet["merged_cells"] is not None else "",
                    embedded_images=sheet["embedded_images"] if sheet["embedded_images"] is not None else "",
                    status=sheet["status"],
                    notes="；".join(sheet["notes"]) if sheet["notes"] else "",
                )
            )
    return "\n".join(lines)


def build_report(
    root: Path,
    workbook_scans: list[dict[str, Any]],
    records: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    column_mapping: dict[str, str],
) -> str:
    sheet_total = sum(workbook["sheet_count"] for workbook in workbook_scans)
    record_counter_by_sheet = Counter((record["source_file"], record["source_sheet"]) for record in records)
    field_counter = Counter()
    for record in records:
        for field_name in PRIORITY_FIELDS + ADDITIONAL_FIELDS:
            if has_meaningful_value(record.get(field_name)):
                field_counter[field_name] += 1

    mapping_groups: dict[str, list[str]] = defaultdict(list)
    for raw_name, target_name in sorted(column_mapping.items()):
        mapping_groups[target_name].append(raw_name)

    abnormal_lines = []
    for issue in issues:
        abnormal_lines.append(
            f"- `{issue['status']}` {issue['file']} / {issue['sheet']}: {issue['reason']}"
        )
    if not abnormal_lines:
        abnormal_lines.append("- 无明显结构异常 sheet。")

    mapping_lines = []
    for target_name, raw_names in sorted(mapping_groups.items()):
        preview = "、".join(raw_names[:6])
        if len(raw_names) > 6:
            preview += f" 等 {len(raw_names)} 个原始列名"
        mapping_lines.append(f"- `{target_name}`: {preview}")

    per_sheet_lines = []
    for workbook in workbook_scans:
        for sheet in workbook["sheets"]:
            key = (workbook["file"], sheet["sheet_name"])
            per_sheet_lines.append(
                f"- `{workbook['file']}` / `{sheet['sheet_name']}`: {record_counter_by_sheet.get(key, 0)} 条记录"
            )

    report = f"""# Excel 清洗报告

## 总览

- 扫描目录：`{root}`
- Excel 文件总数：{len(workbook_scans)}
- Sheet 总数：{sheet_total}
- 识别出的产品记录总数：{len(records)}

## 扫描结果

{render_scan_table(workbook_scans)}

## 每个文件 / 每个 Sheet 的记录数

{chr(10).join(per_sheet_lines) if per_sheet_lines else "- 无"}

## 结构异常或建议复核的 Sheet

{chr(10).join(abnormal_lines)}

## 发现的主要字段名

{chr(10).join(f"- `{field}`: {count} 条记录有值" for field, count in field_counter.most_common())}

## 原始列名到统一字段映射

{chr(10).join(mapping_lines)}

## 清洗规则

- 扫描当前目录及子目录下的 `.xlsx`、`.xls`、`.xlsm` 文件，跳过 `~$` 临时文件。
- 使用合并单元格左上角值补全整块区域，避免跨行产品信息丢失。
- 对文本做轻清洗：去首尾空格、压缩重复空白、统一换行。
- 精确移除明显占位值：空字符串、`—`、`--`、`暂无`、`无`、`未填写`、`未标注` 等。
- 自动识别表头行，优先抽取“每一行是一条产品记录”的 sheet。
- 为每条记录保留 `source_file`、`source_sheet`、`source_row` 溯源字段。
- 自动标准化常见字段名；无法确认语义的列保留到 `extra_fields`。
- 对图片相关列和嵌入图片保留 `image_reference` / `has_image` / `image_columns` 标记，不修改原始 Excel。
- 对结构稀疏或表头识别置信度偏低的 sheet 标记为 `review` 或 `abnormal`，避免强行误解析。

## 前 5 条样本预览

{render_sample_table(records, limit=5)}
"""
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Clean and standardize Excel product sheets.")
    parser.add_argument("--root", default=str(ROOT), help="Project root to scan.")
    parser.add_argument("--output", default="output", help="Output directory relative to root.")
    args = parser.parse_args()

    root = Path(args.root).resolve()
    output_dir = (root / args.output).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    excel_files = discover_excel_files(root)
    workbook_scans: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    column_mapping: dict[str, str] = {}

    for workbook_path in excel_files:
        suffix = workbook_path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            workbook_records, workbook_scan, workbook_issues = process_xlsx_workbook(
                workbook_path, root, column_mapping
            )
        else:
            workbook_records, workbook_scan, workbook_issues = process_xls_workbook(
                workbook_path, root, column_mapping
            )
        records.extend(workbook_records)
        workbook_scans.append(workbook_scan)
        issues.extend(workbook_issues)

    records = sorted(records, key=lambda item: (item["source_file"], item["source_sheet"], item["source_row"]))

    products_json_path = output_dir / "products_cleaned.json"
    products_csv_path = output_dir / "products_cleaned.csv"
    products_xlsx_path = output_dir / "products_cleaned.xlsx"
    column_mapping_path = output_dir / "column_mapping.json"
    report_path = output_dir / "clean_report.md"
    check_report_path = output_dir / "products_cleaned_check_report.md"

    with products_json_path.open("w", encoding="utf-8") as file:
        json.dump(records, file, ensure_ascii=False, indent=2)

    csv_rows = [flatten_record_for_csv(record) for record in records]
    write_csv_rows(products_csv_path, csv_rows)
    write_xlsx_rows_if_possible(products_xlsx_path, csv_rows)

    with column_mapping_path.open("w", encoding="utf-8") as file:
        json.dump(dict(sorted(column_mapping.items())), file, ensure_ascii=False, indent=2)

    report = build_report(root, workbook_scans, records, issues, column_mapping)
    report_path.write_text(report, encoding="utf-8")
    check_report_path.write_text(build_products_check_report(records, issues), encoding="utf-8")

    sheet_total = sum(workbook["sheet_count"] for workbook in workbook_scans)
    print(f"Scanned Excel files: {len(excel_files)}")
    print(f"Scanned sheets: {sheet_total}")
    print(f"Extracted product records: {len(records)}")
    print(f"Output directory: {output_dir}")
    print(f"Output XLSX: {products_xlsx_path}")
    print(f"Check report: {check_report_path}")


if __name__ == "__main__":
    main()
