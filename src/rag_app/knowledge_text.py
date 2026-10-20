"""Bounded, non-executing intake parsers and disclosed heuristic redaction."""

from __future__ import annotations

import hashlib
import html
import io
import json
import re
import stat
import zipfile
from pathlib import PurePosixPath
from urllib.parse import unquote
from xml.etree import ElementTree as ET


MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_MEMBERS = 10000
MAX_TEXT_CHARS = 500000
MAX_LINES = 20000
MAX_ROWS = 20000
MAX_CELLS = 100000
MAX_COLUMNS = 256
MAX_PAGES = 500
TEXT_TYPES = {".md", ".txt", ".csv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SECRET_NAME = r"(?:api[_-]?key|access[_-]?(?:token|key(?:[_-]?(?:id|secret))?)|refresh[_-]?token|client[_-]?secret|password|passwd|secret|token|authorization|cookie|set-cookie)"
ASSIGNMENT = re.compile(
    r"(?P<prefix>(?<![\w:])" + SECRET_NAME + r"\b[\"']?\s*[:=]\s*)"
    r"(?:\"(?P<double>[^\"\r\n]+)\"|'(?P<single>[^'\r\n]+)'|(?P<bare>[^\s,;&<>`}\]]+))", re.I,
)
URL = re.compile(r"(?:https?|ftp)://[^\s<>\"'`]+", re.I)
REDACTIONS = [
    ("private_key", re.compile(r"-----BEGIN [^-\n]*PRIVATE KEY-----.*?-----END [^-\n]*PRIVATE KEY-----", re.S)),
    ("api_key", re.compile(r"\b(?:sk-(?:proj-)?[A-Za-z0-9_-]{20,}|gh[pousr]_[A-Za-z0-9]{20,})\b")),
    ("bearer_token", re.compile(r"\bBearer\s+([A-Za-z0-9_.~+/=-]{8,})", re.I)),
    ("platform_account", re.compile(r"\bpdd\d{5,}\b", re.I)),
    ("email", re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")),
    ("phone_candidate", re.compile(r"(?<![A-Za-z0-9])(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")),
    ("identity_candidate", re.compile(r"(?<![A-Za-z0-9])\d{17}[\dXx](?![A-Za-z0-9])")),
]


def credential_name(name: str) -> bool:
    base = PurePosixPath(name.replace("\\", "/")).name.lower()
    return base.startswith(".env") or base in {
        "auth.json", "credentials.json", "cookies.txt", "cookies.json", "id_rsa", "id_ed25519",
    } or base.endswith((".pem", ".key", ".p12", ".pfx"))


def clean_text(text: str) -> tuple[str, dict]:
    """Preserve content/indentation; stable pseudonyms are not a privacy review."""
    counts = {}

    def hide(kind, value):
        if "[已隐藏:" in value:
            return value
        counts[kind] = counts.get(kind, 0) + 1
        canonical = value
        if kind == "phone_candidate":
            canonical = re.sub(r"\D", "", value)[-11:]
        # ponytail: deterministic pseudonyms preserve references, not anonymization;
        # use an access-controlled mapping if resistance to guessing is required.
        digest = hashlib.sha256((kind + "\0" + canonical).encode()).hexdigest()[:12]
        return "[已隐藏:" + kind + ":" + digest + "]"

    def assignments(match):
        group = next(key for key in ("double", "single", "bare") if match.group(key) is not None)
        value = match.group(group)
        if "[已隐藏:" in value:
            return match.group()
        quote = {"double": '"', "single": "'", "bare": ""}[group]
        return match.group("prefix") + quote + hide("credential_assignment", value) + quote

    def scrub(part, in_url=False):
        part = ASSIGNMENT.sub(assignments, part)
        for kind, pattern in REDACTIONS:
            if in_url and kind in {"phone_candidate", "identity_candidate"}:
                continue
            if kind == "bearer_token":
                part = pattern.sub(lambda m: "Bearer " + hide(kind, m.group(1)), part)
            else:
                part = pattern.sub(lambda m: hide(kind, m.group()), part)
        return part

    def scrub_url(value):
        value = re.sub(r"(?<=//)[^/@?#]+@", lambda m: hide("url_credentials", m.group()[:-1]) + "@", value)
        # Inspect decoded query keys without decoding/reformatting ordinary URLs.
        value = re.sub(
            r"([?&#])([^=&#]+)=([^&#]*)",
            lambda m: m.group(1) + m.group(2) + "=" + (
                hide("credential_assignment", unquote(m.group(3)))
                if re.fullmatch(SECRET_NAME, unquote(m.group(2)), re.I) else m.group(3)
            ), value,
        )
        return scrub(value, in_url=True)

    text = text.removeprefix("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    parts, start = [], 0
    for match in URL.finditer(text):
        parts.extend((scrub(text[start:match.start()]), scrub_url(match.group())))
        start = match.end()
    parts.append(scrub(text[start:]))
    return "".join(parts), counts


class _ParseFailure(ValueError):
    """Only fixed, non-sensitive error codes may cross the parser boundary."""


def _issue(issues, code, **location):
    item = {"code": code, **location}
    if item not in issues:
        issues.append(item)


def _office_archive(data, issues):
    archive = zipfile.ZipFile(io.BytesIO(data))
    try:
        infos = archive.infolist()
        if len(infos) > MAX_MEMBERS or sum(item.file_size for item in infos) > MAX_EXPANDED_BYTES:
            raise _ParseFailure("office_expanded_size_limit")
        names = set()
        for item in infos:
            name = item.filename.replace("\\", "/")
            if (name in names or name.startswith(("/", "~")) or ".." in PurePosixPath(name).parts
                    or re.match(r"^[A-Za-z]:", name) or stat.S_ISLNK(item.external_attr >> 16)):
                raise _ParseFailure("unsafe_office_archive")
            names.add(name)
            if item.flag_bits & 1:
                raise _ParseFailure("encrypted_document")
            if "vbaproject" in name.lower():
                _issue(issues, "macros_not_executed")
            if "/embeddings/" in name.lower() or name.lower().endswith(".zip"):
                _issue(issues, "embedded_content_not_parsed")
            if name.lower().endswith((".xml", ".rels")):
                raw = archive.read(item)
                if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", raw.replace(b"\x00", b""), re.I):
                    raise _ParseFailure("unsafe_xml_declaration")
                if name.lower().endswith(".rels"):
                    if any(node.get("TargetMode", "").lower() == "external" for node in ET.fromstring(raw)):
                        _issue(issues, "external_links_not_fetched")
        return archive
    except Exception:
        archive.close()
        raise


def _paragraph(node):
    return "".join(
        item.text or "" if item.tag == W + "t" else "\t" if item.tag == W + "tab" else "\n"
        for item in node.iter() if item.tag in {W + "t", W + "tab", W + "br", W + "cr"}
    )


def _heading_level(node, styles=None):
    properties = node.find(W + "pPr")
    if properties is None:
        return None
    outline = properties.find(W + "outlineLvl")
    if outline is not None and outline.get(W + "val", "").isdigit():
        return min(6, int(outline.get(W + "val")) + 1) if int(outline.get(W + "val")) < 9 else None
    style = properties.find(W + "pStyle")
    name = style.get(W + "val", "") if style is not None else ""
    match = re.fullmatch(r"(?:Heading|标题)\s*([1-6])", name, re.I)
    return int(match.group(1)) if match else (styles or {}).get(name)


def _docx_styles(archive):
    if "word/styles.xml" not in archive.namelist():
        return {}
    styles = {}
    for node in ET.fromstring(archive.read("word/styles.xml")).findall(W + "style"):
        name = node.find(W + "name")
        outline = node.find(W + "pPr/" + W + "outlineLvl")
        match = re.fullmatch(r"(?:heading|标题)\s*([1-6])", name.get(W + "val", ""), re.I) if name is not None else None
        if outline is not None and outline.get(W + "val", "") in set("012345"):
            styles[node.get(W + "styleId")] = int(outline.get(W + "val")) + 1
        elif match:
            styles[node.get(W + "styleId")] = int(match.group(1))
    return styles


def _table_rows(block):
    return [["\n".join(_paragraph(p) for p in cell.iter(W + "p")) for cell in row.findall(W + "tc")]
            for row in block.findall(W + "tr")]


def _paired_steps(left, right):
    """Join only adjacent, explicitly numbered step tables with the same keys."""
    if (len(left) < 2 or len(right) < 2 or not left[0] or not right[0]
            or left[0][0].strip() not in {"步骤", "步骤号"} or right[0][0].strip() not in {"步骤", "步骤号"}
            or not any("验收" in cell for cell in right[0])):
        return None
    keys = [row[0].strip() for row in left[1:] if row]
    other = {row[0].strip(): row for row in right[1:] if row}
    if (not all(re.fullmatch(r"\d+", key) for key in keys) or len(set(keys)) != len(left) - 1
            or len(other) != len(right) - 1 or set(keys) != set(other)
            or any(len(row) != len(left[0]) for row in left) or any(len(row) != len(right[0]) for row in right)):
        return None
    return [left[0] + right[0][1:]] + [row + other[row[0].strip()][1:] for row in left[1:]]


def _markdown_row(cells):
    return "| " + " | ".join(cell.replace("|", "\\|").replace("\n", "<br>") for cell in cells) + " |"


def _empty_docx_auxiliary(archive, name):
    part = PurePosixPath(name).stem
    root = ET.fromstring(archive.read(name))
    if root.tag != W + part or any(text.strip() for text in root.itertext()):
        return False
    if set(root.attrib) - {"{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"}:
        return False
    if part == "comments":
        return len(root) == 0
    # Ignore only standard separator scaffolding; unknown content still needs review.
    for note in root:
        kind = note.get(W + "type")
        if (note.tag != W + part[:-1] or kind not in {"separator", "continuationSeparator"}
                or set(note.attrib) - {W + "type", W + "id"}
                or not re.fullmatch(r"-?\d+", note.get(W + "id", ""))):
            return False
        children = {note.tag: {W + "p"}, W + "p": {W + "pPr", W + "r"},
                    W + "pPr": {W + "spacing"}, W + "r": {W + kind}}
        attributes = {
            W + "p": {W + key for key in ("rsidR", "rsidRDefault", "rsidP", "rsidDel")} |
                     {"{http://schemas.microsoft.com/office/word/2010/wordml}" + key for key in ("paraId", "textId")},
            W + "spacing": {W + key for key in ("before", "after", "line", "lineRule")},
        }
        if (any(child.tag not in children.get(node.tag, set()) for node in note.iter() for child in node)
                or any(set(node.attrib) - attributes.get(node.tag, set()) for node in note.iter() if node is not note)
                or len(list(note.iter(W + kind))) != 1):
            return False
    return True


def _docx(archive, issues):
    root = ET.fromstring(archive.read("word/document.xml"))
    body = root.find(W + "body")
    if body is None:
        raise _ParseFailure("invalid_docx_body")
    if any(node.tag in {W + "drawing", W + "pict", W + "altChunk", W + "object"} for node in body.iter()):
        _issue(issues, "docx_nontext_content_not_parsed")
    if any(not _empty_docx_auxiliary(archive, name) for name in archive.namelist()
           if re.fullmatch(r"word/(?:footnotes|endnotes|comments)\.xml", name)):
        _issue(issues, "docx_auxiliary_parts_not_parsed")
    styles = _docx_styles(archive)
    paragraph, table, skipped = 0, 0, set()
    blocks = list(body)
    for index, block in enumerate(blocks):
        if index in skipped:
            continue
        if block.tag == W + "p":
            paragraph += 1
            level = _heading_level(block, styles)
            text = _paragraph(block)
            yield ("#" * level + " " + text.replace("\n", " ") if level and text.strip() else text), {
                "kind": "heading" if level else "paragraph", "paragraph": paragraph,
            }
        elif block.tag == W + "tbl":
            table += 1
            if block.find(".//" + W + "tbl") is not None:
                _issue(issues, "docx_nested_table_flattened", table=table)
            rows = _table_rows(block)
            next_index = index + 1
            while next_index < len(blocks) and blocks[next_index].tag == W + "p" and not _paragraph(blocks[next_index]).strip():
                next_index += 1
            paired = _paired_steps(rows, _table_rows(blocks[next_index])) if next_index < len(blocks) and blocks[next_index].tag == W + "tbl" else None
            source_tables = [table]
            if paired:
                rows = paired
                source_tables.append(table + 1)
                skipped.add(next_index)
            rectangular = len(rows) > 1 and all(len(row) == len(rows[0]) for row in rows)
            for row_number, cells in enumerate(rows, 1):
                text = _markdown_row(cells) if rectangular else "[表格 %d 行 %d] %s" % (table, row_number, json.dumps(cells, ensure_ascii=False))
                if rectangular and row_number == 1:
                    text += "\n" + _markdown_row(["---"] * len(cells))
                if row_number == len(rows):
                    text += "\n" if rectangular else ""
                yield text, {"kind": "table_row", "table": table, "row": row_number, "source_tables": source_tables}
            if paired:
                table += 1
        elif block.tag != W + "sectPr":
            _issue(issues, "docx_unsupported_block_not_parsed")
    for name in sorted(archive.namelist()):
        match = re.fullmatch(r"word/(header|footer)\d*\.xml", name)
        if not match:
            continue
        auxiliary = ET.fromstring(archive.read(name))
        if auxiliary.tag not in {W + "hdr", W + "ftr"}:
            _issue(issues, "docx_auxiliary_parts_not_parsed")
            continue
        if any(node.tag in {W + "drawing", W + "pict", W + "object"} for node in auxiliary.iter()):
            _issue(issues, "docx_nontext_content_not_parsed")
        texts = [_paragraph(p) for p in auxiliary.iter(W + "p") if _paragraph(p).strip()]
        if texts:
            yield "\n# " + ("页眉" if match.group(1) == "header" else "页脚") + "\n" + "\n".join(texts), {"kind": match.group(1), "part": name}


def _sop_metadata(archive):
    body = ET.fromstring(archive.read("word/document.xml")).find(W + "body")
    title = " ".join(_paragraph(p) for p in list(body)[:8] if p.tag == W + "p")
    if not re.search(r"岗位.*(?:SOP|标准作业|操作流程)|SOP.*岗位", title, re.I):
        return None
    values = {"所属部门": set(), "岗位名称": set(), "文档名称": set()}
    for block in body.findall(W + "tbl"):
        for row in _table_rows(block):
            for index, cell in enumerate(row[:-1]):
                key = cell.strip().rstrip("：:")
                if key == "所属公司及部门":
                    key = "所属部门"
                if key in values and row[index + 1].strip():
                    values[key].add(row[index + 1].strip())
    if all(len(values[key]) == 1 for key in ("所属部门", "岗位名称")):
        department, role_name = next(iter(values["所属部门"])), next(iter(values["岗位名称"]))
        if all(len(value) <= 120 and "待确认" not in value for value in (department, role_name)):
            metadata = {"department": department, "role_name": role_name}
            if len(values["文档名称"]) == 1:
                document_name = next(iter(values["文档名称"]))
                if len(document_name) <= 240 and "待确认" not in document_name:
                    metadata["document_name"] = document_name
            return metadata
    return None


def docx_html_preview(data: bytes, filename: str = "文档.docx") -> bytes:
    """Text/table preview: no active content, external requests or local files."""
    if not isinstance(data, bytes) or len(data) > MAX_FILE_BYTES or not isinstance(filename, str):
        raise ValueError("invalid_preview_input")
    try:
        fragments, length, issues = [], 0, []
        with _office_archive(data, issues) as archive:
            styles = _docx_styles(archive)
            parts = ["word/document.xml"] + sorted(name for name in archive.namelist() if re.fullmatch(r"word/(?:header|footer)\d*\.xml", name))
            for name in parts:
                root = ET.fromstring(archive.read(name))
                body = root.find(W + "body") if name == "word/document.xml" else root
                if body is None:
                    raise _ParseFailure("invalid_docx_body")
                if name != "word/document.xml":
                    fragments.append("<h2>" + ("页眉" if "/header" in name else "页脚") + "</h2>")
                for block in body:
                    if block.tag == W + "p":
                        text = _paragraph(block)
                        length += len(text)
                        level = _heading_level(block, styles)
                        tag = "h" + str(level) if level else "p"
                        fragments.append("<" + tag + ">" + html.escape(text).replace("\n", "<br>") + "</" + tag + ">")
                    elif block.tag == W + "tbl":
                        fragments.append("<table>")
                        for number, row in enumerate(_table_rows(block)):
                            fragments.append("<tr>")
                            for cell in row:
                                length += len(cell)
                                tag = "th" if number == 0 else "td"
                                fragments.append("<" + tag + ">" + html.escape(cell).replace("\n", "<br>") + "</" + tag + ">")
                            fragments.append("</tr>")
                        fragments.append("</table>")
                    if length > MAX_TEXT_CHARS:
                        raise _ParseFailure("preview_character_limit")
            # Describe omitted rich content using the same parser's fixed codes.
            parsed = _collect(_docx(archive, issues), issues)
            if parsed["issues"]:
                fragments.insert(0, "<p>文本预览未包含或未执行部分内容：" + html.escape(", ".join(item["code"] for item in parsed["issues"])) + "。完整内容请下载原件。</p>")
        return ("<!doctype html><html lang='zh-CN'><meta charset='utf-8'>"
                "<meta name='viewport' content='width=device-width,initial-scale=1'>"
                "<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'\">"
                "<title>" + html.escape(filename) + "</title><style>"
                "body{max-width:1100px;margin:32px auto;padding:0 20px;font:16px/1.65 system-ui,-apple-system,'PingFang SC','Microsoft YaHei',sans-serif;color:#17212b}"
                "table{border-collapse:collapse;width:100%;margin:18px 0;overflow-wrap:anywhere}th,td{border:1px solid #ccd2d8;padding:8px;text-align:left;vertical-align:top}th{background:#f2f5f8}p{white-space:pre-wrap}"
                "</style><body>" + "".join(fragments) + "</body></html>").encode("utf-8")
    except _ParseFailure as exc:
        raise ValueError(str(exc)) from None
    except Exception:
        raise ValueError("preview_failed") from None


def _xlsx(data, issues):
    import openpyxl

    formula_book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached_book = None
    try:
        cached_book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        total_rows, total_cells = 0, 0
        for sheet_number, (sheet, cached) in enumerate(zip(formula_book.worksheets, cached_book.worksheets), 1):
            yield "# 工作表 %d：%s" % (sheet_number, sheet.title), {"kind": "sheet", "sheet_index": sheet_number}
            # Ignore inaccurate dimension metadata; iterate the actual XML rows.
            sheet.reset_dimensions()
            cached.reset_dimensions()
            for row_number, (row, cached_row) in enumerate(zip(sheet.iter_rows(), cached.iter_rows()), 1):
                total_rows += 1
                total_cells += len(row)
                if total_rows > MAX_ROWS or total_cells > MAX_CELLS:
                    _issue(issues, "xlsx_rows_or_cells_limit", sheet_index=sheet_number, row=row_number)
                    return
                if len(row) > MAX_COLUMNS:
                    _issue(issues, "xlsx_columns_limit", sheet_index=sheet_number, row=row_number)
                values = []
                for cell, cached_cell in zip(row[:MAX_COLUMNS], cached_row[:MAX_COLUMNS]):
                    if cell.value is None:
                        continue
                    actual = cell.value
                    if cell.data_type == "f" and not isinstance(actual, str):
                        actual = getattr(actual, "text", None)
                        if actual is None:
                            _issue(issues, "xlsx_formula_type_not_parsed", sheet_index=sheet_number, row=row_number)
                    value = json.dumps(actual, ensure_ascii=False, default=str)
                    if cell.data_type == "f":
                        value = "formula=" + value + "; cached_value=" + json.dumps(cached_cell.value, ensure_ascii=False, default=str)
                        _issue(issues, "xlsx_formulas_not_recalculated")
                        if cached_cell.value is None:
                            _issue(issues, "xlsx_formula_cache_missing", sheet_index=sheet_number, row=row_number)
                    if cell.number_format != "General":
                        value += "; number_format=" + json.dumps(cell.number_format, ensure_ascii=False)
                    values.append(cell.coordinate + "=" + value)
                if values:
                    yield "[表 %d 行 %d] %s" % (sheet_number, row_number, " | ".join(values)), {
                        "kind": "sheet_row", "sheet_index": sheet_number, "row": row_number,
                    }
    finally:
        formula_book.close()
        if cached_book is not None:
            cached_book.close()


def _pdf(data, issues):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise _ParseFailure("encrypted_document")
    root = reader.trailer["/Root"]
    if "/OpenAction" in root or "/AA" in root or "/Names" in root:
        _issue(issues, "pdf_actions_and_attachments_not_executed")
    for number, page in enumerate(reader.pages, 1):
        if number > MAX_PAGES:
            _issue(issues, "pdf_page_limit", limit=MAX_PAGES)
            return
        text = page.extract_text() or ""
        if not text.strip():
            _issue(issues, "needs_ocr", page=number)
            yield "", {"kind": "page", "page": number}
        else:
            yield "[第 %d 页]\n%s" % (number, text), {"kind": "page", "page": number}


def _collect(parts, issues):
    fragments, locations, size, lines = [], [], 0, 0
    try:
        for fragment, location in parts:
            separator = int(bool(fragments))
            if lines >= MAX_LINES or size + separator > MAX_TEXT_CHARS:
                _issue(issues, "text_line_limit" if lines >= MAX_LINES else "text_character_limit",
                       limit=MAX_LINES if lines >= MAX_LINES else MAX_TEXT_CHARS)
                break
            room = max(0, MAX_TEXT_CHARS - size - separator)
            clipped = fragment[:room]
            if len(clipped) < len(fragment):
                _issue(issues, "text_character_limit", limit=MAX_TEXT_CHARS)
            split = clipped.split("\n", max(0, MAX_LINES - lines))
            if len(split) > MAX_LINES - lines:
                clipped = "\n".join(split[:MAX_LINES - lines])
                _issue(issues, "text_line_limit", limit=MAX_LINES)
            fragments.append(clipped)
            locations.append(location)
            size += len(clipped) + separator
            lines += clipped.count("\n") + 1
            if clipped != fragment:
                break
    finally:
        if hasattr(parts, "close"):
            parts.close()
    text = "\n".join(fragments)
    if not text.strip():
        _issue(issues, "no_text_extracted")
    return {"text": text, "locations": locations, "issues": issues}


def parse_text_file(name: str, data: bytes) -> dict:
    """Return unredacted text with locations/issues; caller must run clean_text.

    No execution, external fetches, OCR, decryption, or formula recalculation.
    Truncated output always carries a limit issue; failures never echo input.
    """
    if not isinstance(name, str) or not isinstance(data, bytes):
        raise ValueError("invalid_parser_input")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("file_size_limit")
    if credential_name(name):
        raise ValueError("credential_file_skipped")
    suffix = PurePosixPath(name).suffix.lower()
    issues = []
    try:
        if suffix in TEXT_TYPES:
            if data.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
                raise _ParseFailure("unsupported_text_encoding")
            encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            try:
                text = data.decode(encoding).replace("\r\n", "\n").replace("\r", "\n")
            except UnicodeError:
                if encoding == "utf-16":
                    raise _ParseFailure("invalid_text_encoding") from None
                raise
            if encoding == "utf-16":
                _issue(issues, "encoding_fallback", encoding="utf-16-le" if data.startswith(b"\xff\xfe") else "utf-16-be")
            if "\x00" in text:
                raise _ParseFailure("non_text_content")
            return _collect(iter([(text, {"kind": "lines", "start_line": 1})]), issues)
        if suffix in {".docx", ".xlsx"}:
            with _office_archive(data, issues) as archive:
                result = _collect(_docx(archive, issues) if suffix == ".docx" else _xlsx(data, issues), issues)
                metadata = _sop_metadata(archive) if suffix == ".docx" else None
                if metadata:
                    result["sop_metadata"] = metadata
                return result
        if suffix == ".pdf":
            return _collect(_pdf(data, issues), issues)
        raise _ParseFailure("unsupported_file_type")
    except _ParseFailure as exc:
        raise ValueError(str(exc)) from None
    except UnicodeError:
        raise ValueError("non_utf8_text") from None
    except Exception:
        raise ValueError("parse_failed") from None
